import os
import sys
import glob
import time
import threading
import tkinter as tk
from tkinter import ttk, messagebox
import pandas as pd
import openpyxl

# Importar motor de procesamiento unificado
import matriz_integrada

# =============================================================================
# CONSTANTES DE DISEÑO CORPORATIVO CEMEX
# =============================================================================
CLR_NAVY        = "#002D72"  # Azul CEMEX Principal
CLR_BLUE_LIGHT  = "#005A9C"  # Azul Secundario
CLR_BG          = "#F1F5F9"  # Fondo gris claro suave
CLR_WHITE       = "#FFFFFF"
CLR_GREEN       = "#28A745"  # Verde Acción
CLR_GREEN_HOVER = "#218838"
CLR_TEXT        = "#1E293B"
CLR_MUTED       = "#64748B"
CLR_BORDER      = "#CBD5E1"

FONT_TITLE = ("Segoe UI", 12, "bold")
FONT_SUB   = ("Segoe UI", 8, "normal")
FONT_BODY  = ("Segoe UI", 9, "normal")
FONT_BOLD  = ("Segoe UI", 9, "bold")
FONT_BTN   = ("Segoe UI", 10, "bold")

class LanzadorCEMEXApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CEMEX | Sistema Integral de Matriz de Precios y Gobernanza")
        self.root.geometry("580x640")
        self.root.minsize(480, 480)
        self.root.configure(bg=CLR_BG)
        
        # Centrar ventana en pantalla
        self.centrar_ventana(580, 640)
        
        self.base_dir = os.path.dirname(os.path.abspath(__file__))
        self.source_file = self.detectar_archivo_fuente()
        self.all_cedis = self.cargar_lista_cedis()
        
        # Diccionario para almacenar el estado booleano de cada CEDIS
        self.cedis_vars = {}
        for c in self.all_cedis:
            var = tk.BooleanVar(value=False)
            if c in ["D836", "D838"]:
                var.set(True)
            self.cedis_vars[c] = var
            
        self.checkbox_widgets = {}
        self.is_processing = False
        
        self.construir_interfaz()
        self.actualizar_contador()

    def centrar_ventana(self, ancho, alto):
        self.root.update_idletasks()
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        x = max(0, (sw - ancho) // 2)
        y = max(0, (sh - alto) // 2)
        self.root.geometry(f"{ancho}x{alto}+{x}+{y}")

    def detectar_archivo_fuente(self):
        patron = os.path.join(self.base_dir, "*2026*.xlsm")
        files = glob.glob(patron)
        if files:
            return files[0]
        return os.path.join(self.base_dir, "Extracción de Datos Agregados _ Formatos Para Alta de Ruta (ultimo)_2026.xlsm")

    def cargar_lista_cedis(self):
        try:
            if os.path.exists(self.source_file):
                wb = openpyxl.load_workbook(self.source_file, read_only=True, data_only=True)
                ws = wb['PVTA_MAT VK13']
                headers = next(ws.iter_rows(max_row=1, values_only=True))
                c_idx = headers.index('Centro')
                c_set = set()
                for r in ws.iter_rows(min_row=2, values_only=True):
                    val = r[c_idx]
                    if val:
                        c_set.add(str(val).strip().upper())
                wb.close()
                return sorted(list(c_set))
        except Exception:
            pass
        return ['D836', 'D838', 'DW66', 'D285', 'D847', 'D865', 'D881', 'DW74', 'DW75', 'DW88']

    def construir_interfaz(self):
        # 1. HEADER HERO (ARRIBA)
        header_frame = tk.Frame(self.root, bg=CLR_NAVY, padx=14, pady=8)
        header_frame.pack(fill="x", side="top")
        
        lbl_brand = tk.Label(header_frame, text="CEMEX AGREGADOS", font=("Segoe UI", 8, "bold"), fg="#93C5FD", bg=CLR_NAVY)
        lbl_brand.pack(anchor="w")
        
        lbl_title = tk.Label(header_frame, text="Sistema Integral de Matriz y Gobernanza", font=FONT_TITLE, fg=CLR_WHITE, bg=CLR_NAVY)
        lbl_title.pack(anchor="w")
        
        lbl_sub = tk.Label(header_frame, text="Cálculo automático de MOP %, Desglose de Flete y Dashboard Ejecutivo", font=FONT_SUB, fg="#CBD5E1", bg=CLR_NAVY)
        lbl_sub.pack(anchor="w", pady=(1, 0))

        # 2. FOOTER FIJO (SIEMPRE VISIBLE ABAJO)
        footer_frame = tk.Frame(self.root, bg=CLR_BG, padx=14, pady=8)
        footer_frame.pack(fill="x", side="bottom")

        self.progress_bar = ttk.Progressbar(footer_frame, mode="indeterminate")
        self.lbl_status = tk.Label(footer_frame, text="", font=FONT_SUB, fg=CLR_BLUE_LIGHT, bg=CLR_BG)

        self.btn_generar = tk.Button(
            footer_frame,
            text="🚀  GENERAR MATRIZ Y DASHBOARD (EXCEL)",
            font=FONT_BTN,
            bg=CLR_GREEN,
            fg=CLR_WHITE,
            activebackground=CLR_GREEN_HOVER,
            activeforeground=CLR_WHITE,
            relief="flat",
            pady=8,
            cursor="hand2",
            command=self.iniciar_generacion
        )
        self.btn_generar.pack(fill="x")

        # 3. CONTENEDOR PRINCIPAL (CENTRO CON AUTO-AJUSTE)
        main_frame = tk.Frame(self.root, bg=CLR_BG, padx=14, pady=6)
        main_frame.pack(fill="both", expand=True)

        # 4. BUSCADOR Y ACCIONES RÁPIDAS
        lbl_busc = tk.Label(main_frame, text="🔍 Buscar y seleccionar centros (CEDIS):", font=FONT_BOLD, fg=CLR_TEXT, bg=CLR_BG)
        lbl_busc.pack(anchor="w", pady=(0, 2))
        
        search_box = tk.Frame(main_frame, bg=CLR_BG)
        search_box.pack(fill="x", pady=(0, 4))
        
        self.txt_search = ttk.Entry(search_box, font=FONT_BODY)
        self.txt_search.pack(side="left", fill="x", expand=True, padx=(0, 6))
        self.txt_search.bind("<KeyRelease>", self.filtrar_lista)

        btn_all = tk.Button(search_box, text="☑ Todos", font=FONT_SUB, bg=CLR_WHITE, fg=CLR_TEXT, relief="groove", command=self.marcar_todos, padx=6, pady=1, cursor="hand2")
        btn_all.pack(side="left", padx=2)
        
        btn_none = tk.Button(search_box, text="☐ Ninguno", font=FONT_SUB, bg=CLR_WHITE, fg=CLR_TEXT, relief="groove", command=self.desmarcar_todos, padx=6, pady=1, cursor="hand2")
        btn_none.pack(side="left", padx=2)
        
        btn_demo = tk.Button(search_box, text="🎯 Demo", font=FONT_SUB, bg=CLR_WHITE, fg=CLR_BLUE_LIGHT, relief="groove", command=self.marcar_demo, padx=6, pady=1, cursor="hand2")
        btn_demo.pack(side="left", padx=2)

        # 5. LISTA SCROLLABLE DE CHECKBOXES
        list_card = tk.Frame(main_frame, bg=CLR_WHITE, bd=1, relief="solid", highlightthickness=0)
        list_card.pack(fill="both", expand=True, pady=(2, 4))
        
        self.canvas = tk.Canvas(list_card, bg=CLR_WHITE, height=130, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(list_card, orient="vertical", command=self.canvas.yview)
        self.scrollable_frame = tk.Frame(self.canvas, bg=CLR_WHITE)

        self.scrollable_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )

        self.canvas_window = self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.canvas.bind('<Configure>', self.ajustar_ancho_canvas)
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")
        self.canvas.bind_all("<MouseWheel>", self.on_mousewheel)

        self.dibujar_checkboxes(self.all_cedis)

        # 6. CONTADOR DE SELECCIÓN
        self.lbl_count = tk.Label(main_frame, text="", font=FONT_SUB, fg=CLR_MUTED, bg=CLR_BG)
        self.lbl_count.pack(anchor="w", pady=(0, 4))

        # 7. FILTRO DE CONTRATOS DE COMPRA (TRAOPE) EN FORMATO COMPACTO HORIZONTAL
        frame_opc = tk.LabelFrame(main_frame, text="  Vigencia Contratos Compra (TRAOPE):  ", font=FONT_BOLD, fg=CLR_TEXT, bg=CLR_BG, padx=8, pady=2, relief="groove")
        frame_opc.pack(fill="x", pady=(0, 4))
        
        self.filtro_traope_var = tk.StringVar(value="2024")
        
        rb_2024 = tk.Radiobutton(
            frame_opc,
            text="Vigencia >= 2024 [Recomendado]",
            variable=self.filtro_traope_var,
            value="2024",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_2024.pack(side="left", padx=(0, 16), pady=1)
        
        rb_2029 = tk.Radiobutton(
            frame_opc,
            text="Vigencia >= 2029 [Modo Estricto]",
            variable=self.filtro_traope_var,
            value="2029",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_2029.pack(side="left", padx=0, pady=1)

        # 8. ORGANIZACIÓN DE HOJAS / PESTAÑAS EN EXCEL (CUADRÍCULA 2x2 COMPACTA)
        frame_fmt = tk.LabelFrame(main_frame, text="  Organización de Pestañas en Excel:  ", font=FONT_BOLD, fg=CLR_TEXT, bg=CLR_BG, padx=8, pady=2, relief="groove")
        frame_fmt.pack(fill="x", pady=(0, 2))
        
        self.formato_hojas_var = tk.StringVar(value="por_cedis")
        
        grid_fmt = tk.Frame(frame_fmt, bg=CLR_BG)
        grid_fmt.pack(fill="x", expand=True)

        rb_fmt1 = tk.Radiobutton(
            grid_fmt,
            text="📑 Pestaña x CEDIS [Recomendado]",
            variable=self.formato_hojas_var,
            value="por_cedis",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_fmt1.grid(row=0, column=0, sticky="w", padx=(0, 10), pady=1)
        
        rb_fmt2 = tk.Radiobutton(
            grid_fmt,
            text="📊 1 Sola Hoja Consolidada",
            variable=self.formato_hojas_var,
            value="consolidado",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_fmt2.grid(row=0, column=1, sticky="w", pady=1)

        rb_fmt3 = tk.Radiobutton(
            grid_fmt,
            text="🏢 Pestaña x Sociedad (7100 / 7180)",
            variable=self.formato_hojas_var,
            value="por_sociedad",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_fmt3.grid(row=1, column=0, sticky="w", padx=(0, 10), pady=1)

        rb_fmt4 = tk.Radiobutton(
            grid_fmt,
            text="🌟 Híbrido (Global + x CEDIS)",
            variable=self.formato_hojas_var,
            value="hibrido",
            font=FONT_BODY,
            bg=CLR_BG,
            fg=CLR_TEXT,
            activebackground=CLR_BG,
            selectcolor=CLR_WHITE,
            cursor="hand2"
        )
        rb_fmt4.grid(row=1, column=1, sticky="w", pady=1)

    def ajustar_ancho_canvas(self, event):
        self.canvas.itemconfig(self.canvas_window, width=event.width)

    def on_mousewheel(self, event):
        if self.canvas.winfo_exists():
            self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    def dibujar_checkboxes(self, lista_cedis):
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        self.checkbox_widgets.clear()

        for c in lista_cedis:
            var = self.cedis_vars[c]
            cb = tk.Checkbutton(
                self.scrollable_frame,
                text=f"  CEDIS {c}",
                variable=var,
                font=FONT_BODY,
                bg=CLR_WHITE,
                fg=CLR_TEXT,
                activebackground=CLR_WHITE,
                selectcolor=CLR_WHITE,
                anchor="w",
                padx=8,
                pady=2,
                command=self.actualizar_contador
            )
            cb.pack(fill="x", anchor="w")
            self.checkbox_widgets[c] = cb

    def filtrar_lista(self, event=None):
        query = self.txt_search.get().strip().upper()
        if not query:
            lista_filtrada = self.all_cedis
        else:
            lista_filtrada = [c for c in self.all_cedis if query in c]
        self.dibujar_checkboxes(lista_filtrada)

    def marcar_todos(self):
        for var in self.cedis_vars.values():
            var.set(True)
        self.actualizar_contador()

    def desmarcar_todos(self):
        for var in self.cedis_vars.values():
            var.set(False)
        self.actualizar_contador()

    def marcar_demo(self):
        for c, var in self.cedis_vars.items():
            if c in ["D836", "D838"]:
                var.set(True)
            else:
                var.set(False)
        self.actualizar_contador()

    def actualizar_contador(self):
        total_marcados = sum(1 for var in self.cedis_vars.values() if var.get())
        total_total = len(self.all_cedis)
        if total_marcados == total_total:
            self.lbl_count.config(text=f"✔ Se procesarán TODOS los {total_total} CEDIS del país")
        else:
            self.lbl_count.config(text=f"Centros seleccionados: {total_marcados} de {total_total}")

    def iniciar_generacion(self):
        if self.is_processing:
            return
            
        seleccionados = [c for c, var in self.cedis_vars.items() if var.get()]
        if not seleccionados:
            messagebox.showwarning(
                "Selección Requerida",
                "Por favor seleccione al menos un CEDIS de la lista o presione 'Marcar Todos'."
            )
            return

        if len(seleccionados) == len(self.all_cedis):
            arg_cedis = None  # Indica procesar todos
            txt_cedis = "TODOS LOS CEDIS"
        else:
            arg_cedis = seleccionados
            txt_cedis = ", ".join(seleccionados[:4]) + (f" (+{len(seleccionados)-4} más)" if len(seleccionados)>4 else "")

        # Bloquear UI y mostrar barra de progreso
        self.is_processing = True
        self.btn_generar.config(state="disabled", text="⏳  PROCESANDO DATOS...", bg=CLR_MUTED)
        self.progress_bar.pack(fill="x", pady=(0, 4), before=self.btn_generar)
        self.progress_bar.start(10)
        self.lbl_status.pack(pady=(0, 4), before=self.btn_generar)
        self.lbl_status.config(text=f"Procesando {txt_cedis}...")

        # Ejecutar procesamiento en un hilo secundario para no congelar la UI
        hilo = threading.Thread(target=self.ejecutar_motor, args=(arg_cedis,), daemon=True)
        hilo.start()

    def ejecutar_motor(self, arg_cedis):
        out_dir = os.path.join(self.base_dir, "_salidas_integradas")
        os.makedirs(out_dir, exist_ok=True)
        
        filtro_sel = self.filtro_traope_var.get()
        fmt_sel = self.formato_hojas_var.get()
        
        try:
            # 1. Cargar datos maestros (usa caché rápido)
            data_raw = matriz_integrada.cargar_excel(self.source_file, force_refresh=False)
            
            # 2. Procesar cruce, MOP puro, gobernanza y validaciones con la vigencia seleccionada
            df_matriz, df_mp, df_flete, df_traope, df_contratos = matriz_integrada.procesar_datos(
                data_raw, cedis=arg_cedis, modo_a=True, filtro_traope=filtro_sel
            )
            
            # 3. Construir libro final con Dashboard y Matriz según formato de hojas
            archivo_generado = matriz_integrada.escribir_excel(
                df_matriz, arg_cedis, out_dir, [df_mp, df_flete, df_traope, df_contratos],
                filtro_traope=filtro_sel, formato_hojas=fmt_sel
            )
            
            # Notificar éxito a la UI principal
            self.root.after(0, self.finalizar_exito, archivo_generado)
            
        except Exception as e:
            self.root.after(0, self.finalizar_error, str(e))

    def finalizar_exito(self, archivo_generado):
        self.progreso_terminado()
        
        # Abrir automáticamente el archivo generado en Microsoft Excel
        if archivo_generado and os.path.exists(archivo_generado):
            try:
                os.startfile(archivo_generado)
            except Exception:
                pass
            
            nombre_arc = os.path.basename(archivo_generado)
            messagebox.showinfo(
                "¡Proceso Completado!",
                f"Matriz y Dashboard generados exitosamente.\n\n"
                f"Archivo: {nombre_arc}\n\n"
                f"Se ha abierto el reporte en Microsoft Excel."
            )
        else:
            messagebox.showinfo(
                "Proceso Terminado",
                "Proceso completado. Los reportes se guardaron en la carpeta _salidas_integradas."
            )

    def finalizar_error(self, mensaje_error):
        self.progreso_terminado()
        messagebox.showerror(
            "Error en Generación",
            f"Ocurrió un inconveniente durante el cálculo:\n\n{mensaje_error}"
        )

    def progreso_terminado(self):
        self.is_processing = False
        self.progress_bar.stop()
        self.progress_bar.pack_forget()
        self.lbl_status.pack_forget()
        self.btn_generar.config(
            state="normal",
            text="🚀  GENERAR MATRIZ Y DASHBOARD (EXCEL)",
            bg=CLR_GREEN
        )

# =============================================================================
# ENTRADA PRINCIPAL
# =============================================================================
if __name__ == "__main__":
    root = tk.Tk()
    app = LanzadorCEMEXApp(root)
    root.mainloop()

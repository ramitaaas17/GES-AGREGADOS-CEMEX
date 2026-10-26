# -*- coding: utf-8 -*-
"""
Módulo extractor y serializador de los Catálogos Maestros de la Jefa:
- 'Listado CEDIS' -> Mapeo Centro (Dxxx / DWxx) a Nombre Planta VD03.
- 'Clientes' -> Mapeo Destino / Frente Clave a Nombre Frente / Destinatario.
"""
import os
import pickle
import openpyxl
import pandas as pd

import sys

RUTA_JEFA = r"C:\Users\E-FACORONAV\Downloads\Matriz_Precios_TODOS_Master_Consolidado_C2024_20261010_132553 (1).xlsx"
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PKL = os.path.join(BASE_DIR, "catalogos_maestros_jefa.pkl")


def normalizar_clave(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.upper()


def compilar_catalogos(ruta_excel=RUTA_JEFA, out_pkl=OUT_PKL):
    print(f"Cargando catálogos desde: {ruta_excel}...")
    if not os.path.exists(ruta_excel):
        raise FileNotFoundError(f"No se encontró el archivo: {ruta_excel}")

    wb = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)

    # 1. LISTADO CEDIS
    map_cedis = {}
    if 'Listado CEDIS' in wb.sheetnames:
        ws_cedis = wb['Listado CEDIS']
        for r in ws_cedis.iter_rows(min_row=2, values_only=True):
            if r and len(r) >= 2 and r[0] is not None:
                c_clave = normalizar_clave(r[0])
                c_nom = str(r[1] or "").strip()
                if c_clave and c_nom:
                    map_cedis[c_clave] = c_nom
        print(f"  [OK] Listado CEDIS: {len(map_cedis)} centros indexados.")
    else:
        print("  [WARN] No se encontró la pestaña 'Listado CEDIS'.")

    # 2. CLIENTES
    map_destinos = {}
    if 'Clientes' in wb.sheetnames:
        ws_cli = wb['Clientes']
        headers = [str(c or "").replace('\u202a', '').replace('\u202c', '').strip().upper() for c in next(ws_cli.iter_rows(min_row=1, max_row=1, values_only=True))]
        
        # Localizar índices de columnas
        idx_frente_c = headers.index('FRENTE CLAVE') if 'FRENTE CLAVE' in headers else 12
        idx_frente_n = headers.index('FRENTE') if 'FRENTE' in headers else 13
        idx_dest_c = headers.index('DESTINATARIO DE MERCANCÍA CLAVE') if 'DESTINATARIO DE MERCANCÍA CLAVE' in headers else (headers.index('DESTINATARIO DE MERCANCIA CLAVE') if 'DESTINATARIO DE MERCANCIA CLAVE' in headers else 8)
        idx_dest_n = headers.index('DESTINATARIO DE MERCANCÍA') if 'DESTINATARIO DE MERCANCÍA' in headers else (headers.index('DESTINATARIO DE MERCANCIA') if 'DESTINATARIO DE MERCANCIA' in headers else 9)

        for r in ws_cli.iter_rows(min_row=2, values_only=True):
            if not r:
                continue
            # Frente
            if len(r) > idx_frente_c and r[idx_frente_c] is not None:
                f_c = normalizar_clave(r[idx_frente_c])
                f_n = str(r[idx_frente_n] or "").strip() if len(r) > idx_frente_n and r[idx_frente_n] is not None else ""
                if f_c and f_n and f_c not in map_destinos:
                    map_destinos[f_c] = f_n

            # Destinatario mercancía
            if len(r) > idx_dest_c and r[idx_dest_c] is not None:
                d_c = normalizar_clave(r[idx_dest_c])
                d_n = str(r[idx_dest_n] or "").strip() if len(r) > idx_dest_n and r[idx_dest_n] is not None else ""
                if d_c and d_n and d_c not in map_destinos:
                    map_destinos[d_c] = d_n

        print(f"  [OK] Clientes: {len(map_destinos)} destinos/frentes indexados.")
    else:
        print("  [WARN] No se encontró la pestaña 'Clientes'.")

    wb.close()

    catalogos = {
        "cedis": map_cedis,
        "destinos": map_destinos,
    }

    with open(out_pkl, "wb") as f_out:
        pickle.dump(catalogos, f_out, protocol=pickle.HIGHEST_PROTOCOL)

    print(f"Catálogos guardados exitosamente en: {out_pkl}")
    return catalogos


def cargar_catalogos(out_pkl=OUT_PKL):
    if os.path.exists(out_pkl):
        try:
            with open(out_pkl, "rb") as f_in:
                return pickle.load(f_in)
        except Exception as e:
            print(f"Error cargando caché de catálogos: {e}")
    if os.path.exists(RUTA_JEFA):
        return compilar_catalogos(RUTA_JEFA, out_pkl)
    return {"cedis": {}, "destinos": {}}


if __name__ == "__main__":
    compilar_catalogos()

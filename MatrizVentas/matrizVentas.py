#!/usr/bin/env python3
"""
CEMEX Matriz Ventas Automatizacion
====================================
Genera "Base Cedis DW##" a partir de 4 archivos de entrada:
 - MP / VK13 / precios de venta (archivo de precios de materia prima)
 - Fletes / VK13 Fletes          (archivo de fletes)
 - TRAOPE                         (ordenes de compra / condiciones)
 - Materiales                     (catalogo de materiales con PV)


Produce 2 archivos de salida:
 - Base Cedis DW## (Con Formulas).xlsx  → con VLOOKUP/BUSCARV formulas
 - Base Cedis DW##.xlsx                 → valores limpios + validaciones


Uso:
   python3 generar_matriz.py <carpeta>
   python3 generar_matriz.py --mp MP.xlsx --flete Fletes.xlsx --traope TRAOPE.xlsx --materiales MAT.xlsx [--output /ruta] [--cedis DW88]
"""


import argparse
import re
import sys
from pathlib import Path
from datetime import datetime
from typing import Optional, Dict


import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ═══════════════════════════════════════════════════════════
#  GRAMATICAS DE RECONOCIMIENTO DE ARCHIVOS (RegEx)
# ═══════════════════════════════════════════════════════════


FILE_PATTERNS: Dict[str, re.Pattern] = {
   # TRAOPE debe ir primero (mas especifico)
   "traope": re.compile(r"traope", re.IGNORECASE),
   # Fletes: contiene "flete/fletes" o es VK13 con la palabra flete
   "flete": re.compile(r"flet[eo]s?|vk13[_\s-]*flet", re.IGNORECASE),
   # Materiales: contiene "material(es)" pero no es flete ni traope
   "materiales": re.compile(r"material(?:es)?", re.IGNORECASE),
   # MP: contiene "mp", "precios de venta/compra", o VK13 sin flete
   "mp": re.compile(
       r"\bmp\b|precios?\s+(?:de\s+)?(?:venta|compra)|vk13(?![_\s-]*flet)",
       re.IGNORECASE,
   ),
}


CEDIS_PATTERN = re.compile(r"\b(DW\s*\d+)\b", re.IGNORECASE)

# Codigos de "clase de condicion" SAP para el renglon de importe en el formato
# vertical. Varian segun la linea de producto del CEDIS (Sacos/Intergiros usa
# ZMAH/ZMPH, Agregados/terceros usa ZMA6/ZMP1), pero el resto de la estructura
# del archivo es identica.
MP_CONDITION_CODES = ("ZMAH", "ZMA6")
FLETE_CONDITION_CODES = ("ZMPH", "ZMP1")


# Extensiones aceptadas
EXCEL_EXTS = {".xlsx", ".xls", ".xlsm", ".xlsb"}




# ═══════════════════════════════════════════════════════════
#  ESTILOS Y FORMATOS
# ═══════════════════════════════════════════════════════════


def _fill(rgb: str) -> PatternFill:
   return PatternFill(patternType="solid", fgColor=rgb)


def _border(style: str = "thin") -> Border:
   s = Side(border_style=style)
   return Border(left=s, right=s, top=s, bottom=s)


def _apply(cell, font=None, fill=None, alignment=None, border=None, number_format=None):
   if font:
       cell.font = font
   if fill:
       cell.fill = fill
   if alignment:
       cell.alignment = alignment
   if border:
       cell.border = border
   if number_format:
       cell.number_format = number_format


# Paleta de colores
CLR_VENTA_HDR  = "1F5C85"   # azul oscuro  → seccion VENTA
CLR_VENTA_COL  = "BDD7EE"   # azul claro   → columnas de VENTA
CLR_COSTO_HDR  = "7B3F00"   # marron       → seccion COSTO
CLR_COSTO_COL  = "FCE4D6"   # salmon claro → columnas de COSTO
CLR_VAL_COL    = "E2EFDA"   # verde claro  → columnas de validacion
CLR_HEADER     = "D9D9D9"   # gris claro   → resto de encabezados
CLR_WARN       = "FFFF00"   # amarillo     → diferencia > 1 en validacion 2


FONT_HDR  = Font(bold=True, size=10)
FONT_SEC  = Font(bold=True, size=11, color="FFFFFF")
ALIGN_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LFT = Alignment(horizontal="left",   vertical="center")
ALIGN_RGT = Alignment(horizontal="right",  vertical="center")
BORDER    = _border("thin")


FMT_NUM2  = "#,##0.00"
FMT_NUM3  = "#,##0.000"


# Columnas del sheet principal (nombre → numero 1-indexed)
COLS = {
   "CONCAT1":         1,
   "CONCAT2":         2,
   "NO_SF":           3,
   "NOMBRE_SF":       4,
   "CEDIS":           5,
   "DESTINO":         6,
   "COND_EXP":        7,
   "NOMBRE_DESTINO":  8,
   "NO_MATERIAL":     9,
   "NOMBRE_MATERIAL": 10,
   "PV":              11,
   "INICIO_VIG":      12,
   "FIN_VIG":         13,
   # -- VENTA --
   "IMPORTE_MP":      14,
   "IMPORTE_FLETE":   15,
   "UM_VENTA":        16,
   # -- COSTO --
   "IMPORTE_COSTO":   17,
   "UM_COSTO":        18,
   # -- VALIDACIONES --
   "VALIDACION1":     19,
   "VALIDACION2":     20,
}


COL_LABELS = {
   "CONCAT1":         "CONCAT1",
   "CONCAT2":         "CONCAT2",
   "NO_SF":           "No. SF",
   "NOMBRE_SF":       "Nombre SF",
   "CEDIS":           "Cedis",
   "DESTINO":         "Destino",
   "COND_EXP":        "Condicion de Expedicion",
   "NOMBRE_DESTINO":  "Nombre Destino",
   "NO_MATERIAL":     "No. de material",
   "NOMBRE_MATERIAL": "Nombre de material",
   "PV":              "PV",
   "INICIO_VIG":      "Inicio de vigencia",
   "FIN_VIG":         "Fin de vigencia",
   "IMPORTE_MP":      "Importe MP",
   "IMPORTE_FLETE":   "Importe Flete",
   "UM_VENTA":        "UM Venta",
   "IMPORTE_COSTO":   "Importe Costo",
   "UM_COSTO":        "UM Costo",
   "VALIDACION1":     "Validacion 1",
   "VALIDACION2":     "Validacion 2",
}


VENTA_KEYS  = {"IMPORTE_MP", "IMPORTE_FLETE", "UM_VENTA"}
COSTO_KEYS  = {"IMPORTE_COSTO", "UM_COSTO"}
VAL_KEYS    = {"VALIDACION1", "VALIDACION2"}


COL_WIDTHS = {
   "CONCAT1":         34, "CONCAT2":         28,
   "NO_SF":           14, "NOMBRE_SF":       28,
   "CEDIS":            8, "DESTINO":         10,
   "COND_EXP":        13, "NOMBRE_DESTINO":  30,
   "NO_MATERIAL":     14, "NOMBRE_MATERIAL": 40,
   "PV":               8, "INICIO_VIG":      14,
   "FIN_VIG":         14, "IMPORTE_MP":      14,
   "IMPORTE_FLETE":   14, "UM_VENTA":        10,
   "IMPORTE_COSTO":   14, "UM_COSTO":        10,
   "VALIDACION1":     14, "VALIDACION2":     14,
}




# Layout del archivo "limpio" (sin formulas): igual que COLS pero sin
# CONCAT1/CONCAT2 (son solo llaves internas de cruce, no le sirven al usuario
# y en ese archivo no hay formulas VLOOKUP que dependan de ellas).
COLS_CLEAN: Dict[str, int] = {}
for _key in COLS:
   if _key in ("CONCAT1", "CONCAT2"):
       continue
   COLS_CLEAN[_key] = len(COLS_CLEAN) + 1
del _key


def _cl(key: str, cols: Dict[str, int] = COLS) -> str:
   """Letra(s) de columna para una clave de COLS."""
   return get_column_letter(cols[key])




def _ref(key: str, row: int, cols: Dict[str, int] = COLS) -> str:
   """Referencia de celda tipo 'N3'."""
   return f"{_cl(key, cols)}{row}"




# ═══════════════════════════════════════════════════════════
#  DETECCION DE ARCHIVOS
# ═══════════════════════════════════════════════════════════


def identify_file(path: Path) -> Optional[str]:
   """Identifica el tipo de archivo por su nombre usando gramaticas regex."""
   name = path.stem
   # Orden importa: traope y materiales antes que mp (evita falsos positivos)
   for ftype in ("traope", "materiales", "flete", "mp"):
       if FILE_PATTERNS[ftype].search(name):
           return ftype
   return None




def find_files_in_folder(folder: Path) -> Dict[str, Path]:
   """Escanea una carpeta y clasifica los archivos Excel."""
   found: Dict[str, Path] = {}
   for path in sorted(folder.iterdir()):
       if path.suffix.lower() not in EXCEL_EXTS:
           continue
       if path.name.startswith("~") or path.name.startswith("."):
           continue
       ftype = identify_file(path)
       if ftype and ftype not in found:
           found[ftype] = path
   return found




def extract_cedis(files: Dict[str, Path]) -> str:
   """Extrae el codigo CEDIS (ej. DW88) de los nombres de archivo."""
   for path in files.values():
       m = CEDIS_PATTERN.search(path.stem)
       if m:
           return m.group(1).upper().replace(" ", "")
   return "DW00"




# ═══════════════════════════════════════════════════════════
#  UTILIDADES DE LECTURA
# ═══════════════════════════════════════════════════════════


def read_excel_safe(path: Path, **kwargs) -> pd.DataFrame:
   """Lee un archivo Excel probando motores en orden."""
   errs = []
   for engine in (None, "openpyxl", "xlrd"):
       try:
           kw = dict(kwargs)
           if engine:
               kw["engine"] = engine
           return pd.read_excel(path, header=None, **kw)
       except Exception as e:
           errs.append(f"{engine}: {e}")
   raise ValueError(f"No se pudo leer '{path.name}': {'; '.join(errs)}")




def str_val(v) -> str:
   """Convierte un valor a string limpio. Los floats enteros se convierten a int."""
   if v is None or (not isinstance(v, str) and pd.isna(v)):
       return ""
   if isinstance(v, float) and v == int(v):
       return str(int(v))
   return str(v).strip()




def date_str(v) -> str:
   """Formatea fechas como DD.MM.YYYY."""
   if v is None or (not isinstance(v, str) and pd.isna(v)):
       return ""
   if isinstance(v, (pd.Timestamp, datetime)):
       return v.strftime("%d.%m.%Y")
   s = str(v).strip()
   # Si ya tiene formato DD.MM.YYYY lo devolvemos tal cual
   if re.match(r"\d{2}\.\d{2}\.\d{4}", s):
       return s
   return s




def find_data_start_row(df: pd.DataFrame, min_non_null: int = 5) -> int:
   """Devuelve el indice de la primera fila con suficientes valores reales."""
   for i, row in df.iterrows():
       count = sum(1 for v in row if not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip())
       if count >= min_non_null:
           return i
   return 0




def find_traope_header_row(df: pd.DataFrame) -> int:
   """Localiza la fila de encabezados del TRAOPE (contiene 'Ship From' u otras claves)."""
   signatures = {"Doc. Compras", "Ship From", "Nombre SF", "Material", "Centro"}
   for i, row in df.iterrows():
       vals = {str(v).strip() for v in row if not (v is None or (not isinstance(v, str) and pd.isna(v)))}
       if len(vals & signatures) >= 2:
           return i
   return 0




# ═══════════════════════════════════════════════════════════
#  PARSEO DE CADA TIPO DE ARCHIVO
# ═══════════════════════════════════════════════════════════


def parse_mp(path: Path) -> pd.DataFrame:
   """
   Parsea el archivo de MP / VK13 / precios de venta.


   Soporta dos formatos:
     - Columnar (formato antiguo): una fila por registro, columnas fijas.
     - Vertical apilado (nuevo formato): ~18 filas por registro con etiquetas en col0/col1.


   Estructura columnar (sin encabezado):
     Col 0 : No. SF
     Col 1 : Cedis / Centro
     Col 2 : Destino / Destinatario
     Col 3 : No. Material
     Col 4 : BLANCO
     Col 5 : Nombre Material
     Col 6 : BLANCO
     Col 7 : Importe MP
     Col 8 : Moneda
     Col 9 : Cantidad
     Col 10: UM (unidad de medida)
     Col 11: C (sin nombre)
     Col 12: BLANCO
     Col 13: Inicio vigencia
     Col 14: Fin vigencia
   """
   raw = read_excel_safe(path)
   if is_vertical_format(raw):
       df = parse_mp_vertical(raw)
       df.attrs["is_vertical"] = True
       return df


   start = find_data_start_row(raw)
   rows = []
   for _, row in raw.iloc[start:].iterrows():
       sf = str_val(row.iloc[0])
       if not sf or not re.match(r"^\d", sf):
           continue
       rows.append({
           "sf":              sf,
           "cedis":           str_val(row.iloc[1]),
           "destino":         str_val(row.iloc[2]),
           "nombre_destino":  "",
           "material":        str_val(row.iloc[3]),
           "nombre_material": str_val(row.iloc[5]) if len(row) > 5 else "",
           "importe_mp":      row.iloc[7] if len(row) > 7 else None,
           "moneda":          str_val(row.iloc[8]) if len(row) > 8 else "",
           "cantidad":        row.iloc[9] if len(row) > 9 else None,
           "um":              str_val(row.iloc[10]) if len(row) > 10 else "",
           "sin_nombre":      str_val(row.iloc[11]) if len(row) > 11 else "",
           "inicio_vig":      date_str(row.iloc[13]) if len(row) > 13 else "",
           "fin_vig":         date_str(row.iloc[14]) if len(row) > 14 else "",
       })
   df = pd.DataFrame(rows)
   df.attrs["is_vertical"] = False
   print(f"   MP: {len(df)} registros")
   return df




def parse_flete(path: Path) -> pd.DataFrame:
   """
   Parsea el archivo de Fletes.


   Soporta dos formatos:
     - Columnar (formato antiguo): una fila por registro, columnas fijas.
     - Vertical apilado (nuevo formato): ~19 filas por registro con etiquetas en col0/col1.


   Estructura columnar (sin encabezado):
     Col 0 : No. SF
     Col 1 : Cedis / Centro
     Col 2 : Destino / Destinatario
     Col 3 : Condicion de Expedicion (1 o 4)
     Col 4 : No. Material
     Col 5 : BLANCO
     Col 6 : Nombre Material
     Col 7 : BLANCO
     Col 8 : Importe Flete
     Col 9 : Moneda
     Col 10: Cantidad
     Col 11: UM
     Col 12: C (sin nombre)
     Col 13: BLANCO
     Col 14: Inicio vigencia
     Col 15: Fin vigencia
   """
   raw = read_excel_safe(path)
   if is_vertical_format(raw):
       df = parse_flete_vertical(raw)
       df.attrs["is_vertical"] = True
       return df


   start = find_data_start_row(raw)
   rows = []
   for _, row in raw.iloc[start:].iterrows():
       sf = str_val(row.iloc[0])
       if not sf or not re.match(r"^\d", sf):
           continue
       rows.append({
           "sf":             sf,
           "cedis":          str_val(row.iloc[1]),
           "destino":        str_val(row.iloc[2]),
           "cond_exp":       str_val(row.iloc[3]),
           "material":       str_val(row.iloc[4]),
           "nombre_material": str_val(row.iloc[6]) if len(row) > 6 else "",
           "importe_flete":  row.iloc[8] if len(row) > 8 else None,
           "moneda":         str_val(row.iloc[9]) if len(row) > 9 else "",
           "cantidad":       row.iloc[10] if len(row) > 10 else None,
           "um":             str_val(row.iloc[11]) if len(row) > 11 else "",
           "sin_nombre":     str_val(row.iloc[12]) if len(row) > 12 else "",
           "inicio_vig":     date_str(row.iloc[14]) if len(row) > 14 else "",
           "fin_vig":        date_str(row.iloc[15]) if len(row) > 15 else "",
       })
   df = pd.DataFrame(rows)
   df.attrs["is_vertical"] = False
   print(f"   Flete: {len(df)} registros")
   return df




# ═══════════════════════════════════════════════════════════
#  FORMATO VERTICAL (nuevo formato MP y Flete)
# ═══════════════════════════════════════════════════════════


def is_vertical_format(df: pd.DataFrame) -> bool:
   """Detecta si el archivo tiene el formato vertical apilado (nuevo formato)."""
   for i in range(min(5, len(df))):
       try:
           col0 = str(df.iloc[i, 0]).strip()
       except Exception:
           continue
       if "Clave del registro" in col0 or col0 == "Acreedor":
           return True
   return False




def parse_mp_vertical(raw: pd.DataFrame) -> pd.DataFrame:
   """
   Parsea archivo MP en formato vertical apilado (nuevo formato).


   Cada registro (~18 filas) esta separado por una fila con:
     col0 = "Clave del registro de condicion"


   Mapeo de campos:
     col0="Acreedor"                   → sf = col2
     col0 starts with "Centro"         → cedis = col2
     col0 contains "Destinatario"      → destino = col2
     col0 starts with "Material"       → material = col2, nombre_material = col5
     col0 contains "validez"           → inicio_vig = col2, fin_vig = col5
     col1 contains ZMAH/ZMA6            → importe_mp=col4, moneda=col6, cantidad=col8, um=col9
   """
   rows = []
   current: dict = {}


   def sv(row, idx):
       return str_val(row.iloc[idx]) if len(row) > idx else ""


   def dv(row, idx):
       return date_str(row.iloc[idx]) if len(row) > idx else ""


   def nv(row, idx):
       return row.iloc[idx] if len(row) > idx else None


   for _, row in raw.iterrows():
       col0 = sv(row, 0)
       col1 = sv(row, 1)


       if "Clave del registro" in col0:
           if current.get("sf"):
               rows.append(current)
           current = {}
           continue


       if col0 == "Acreedor":
           current["sf"] = sv(row, 2)
       elif col0.startswith("Centro"):
           current["cedis"] = sv(row, 2)
       elif "Destinatario" in col0:
           current["destino"] = sv(row, 2)
           current["nombre_destino"] = sv(row, 5)
       elif col0.startswith("Material") and "material" not in current:
           current["material"] = sv(row, 2)
           current["nombre_material"] = sv(row, 5)
       elif "validez" in col0.lower():
           current["inicio_vig"] = dv(row, 2)
           current["fin_vig"]    = dv(row, 5)
       elif any(code in col1 for code in MP_CONDITION_CODES):
           current["importe_mp"] = nv(row, 4)
           current["moneda"]     = sv(row, 6)
           current["cantidad"]   = nv(row, 8)
           current["um"]         = sv(row, 9)
           current["sin_nombre"] = "C"


   if current.get("sf"):
       rows.append(current)


   defaults = {
       "sf": "", "cedis": "", "destino": "", "nombre_destino": "", "material": "",
       "nombre_material": "", "importe_mp": None, "moneda": "",
       "cantidad": None, "um": "", "sin_nombre": "", "inicio_vig": "", "fin_vig": "",
   }
   for r in rows:
       for k, v in defaults.items():
           r.setdefault(k, v)


   result = pd.DataFrame(rows)
   print(f"   MP (vertical): {len(result)} registros")
   return result




def parse_flete_vertical(raw: pd.DataFrame) -> pd.DataFrame:
   """
   Parsea archivo Flete en formato vertical apilado (nuevo formato).


   Igual que MP vertical pero ademas captura:
     col0 contains "expedici"  → cond_exp = col2
     col1 contains ZMPH/ZMP1  → importe_flete=col4, moneda=col6, cantidad=col8, um=col9
   """
   rows = []
   current: dict = {}


   def sv(row, idx):
       return str_val(row.iloc[idx]) if len(row) > idx else ""


   def dv(row, idx):
       return date_str(row.iloc[idx]) if len(row) > idx else ""


   def nv(row, idx):
       return row.iloc[idx] if len(row) > idx else None


   for _, row in raw.iterrows():
       col0 = sv(row, 0)
       col1 = sv(row, 1)


       if "Clave del registro" in col0:
           if current.get("sf"):
               rows.append(current)
           current = {}
           continue


       if col0 == "Acreedor":
           current["sf"] = sv(row, 2)
       elif col0.startswith("Centro"):
           current["cedis"] = sv(row, 2)
       elif "Destinatario" in col0:
           current["destino"] = sv(row, 2)
           current["nombre_destino"] = sv(row, 5)
       elif col0.startswith("Material") and "material" not in current:
           current["material"] = sv(row, 2)
           current["nombre_material"] = sv(row, 5)
       elif "expedici" in col0.lower():
           current["cond_exp"] = sv(row, 2)
       elif "validez" in col0.lower():
           current["inicio_vig"] = dv(row, 2)
           current["fin_vig"]    = dv(row, 5)
       elif any(code in col1 for code in FLETE_CONDITION_CODES):
           current["importe_flete"] = nv(row, 4)
           current["moneda"]        = sv(row, 6)
           current["cantidad"]      = nv(row, 8)
           current["um"]            = sv(row, 9)
           current["sin_nombre"]    = "C"


   if current.get("sf"):
       rows.append(current)


   defaults = {
       "sf": "", "cedis": "", "destino": "", "nombre_destino": "", "material": "",
       "nombre_material": "", "cond_exp": "", "importe_flete": None,
       "moneda": "", "cantidad": None, "um": "", "sin_nombre": "",
       "inicio_vig": "", "fin_vig": "",
   }
   for r in rows:
       for k, v in defaults.items():
           r.setdefault(k, v)


   result = pd.DataFrame(rows)
   print(f"   Flete (vertical): {len(result)} registros")
   return result




def parse_traope(path: Path) -> pd.DataFrame:
   """
   Parsea el archivo TRAOPE.


   Soporta dos variantes:
     - Encabezados en fila 0 (ej. DW63)
     - Fila de titulo + blancos + encabezados en fila 3 (ej. DW88)


   Devuelve DataFrame con columnas nombradas segun encabezados originales,
   mas 'concat1' calculado = ShipFrom + Centro + Destino + Material.
   """
   raw = read_excel_safe(path)
   hdr_idx = find_traope_header_row(raw)


   # Detectar si los datos de encabezado inician en col 0 o col 1
   hdr_series = raw.iloc[hdr_idx]
   col_offset = 0
   if pd.isna(hdr_series.iloc[0]) or str(hdr_series.iloc[0]).strip() == "":
       col_offset = 1


   # Mapear nombre_columna -> indice en raw
   col_map: Dict[str, int] = {}
   for ci, v in enumerate(hdr_series):
       if ci < col_offset:
           continue
       if not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip():
           col_map[str(v).strip()] = ci


   # Encontrar columnas clave por nombre flexible
   def find_col(*names) -> Optional[int]:
       for n in names:
           for k, ci in col_map.items():
               if n.lower() in k.lower():
                   return ci
       return None


   sf_ci          = find_col("Ship From")
   nombre_sf_ci   = find_col("Nombre SF")
   mat_ci         = find_col("Material")
   centro_ci      = find_col("Centro")
   destino_ci     = find_col("Destino")
   precio_ci      = find_col("Precio neto pedido", "Precio neto")
   um_precio_ci   = find_col("UM Precio Pedido")


   # Determinar inicio de datos
   data_start = hdr_idx + 1
   for i in range(data_start, min(data_start + 5, len(raw))):
       non_null = sum(1 for v in raw.iloc[i] if not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip())
       if non_null >= 5:
           data_start = i
           break


   rows = []
   for _, row in raw.iloc[data_start:].iterrows():
       if sf_ci is None:
           continue
       sf = str_val(row.iloc[sf_ci])
       if not sf or not re.match(r"^\d", sf):
           continue


       mat      = str_val(row.iloc[mat_ci])     if mat_ci     is not None else ""
       centro   = str_val(row.iloc[centro_ci])  if centro_ci  is not None else ""
       destino  = str_val(row.iloc[destino_ci]) if destino_ci is not None else ""
       precio_v = row.iloc[precio_ci]            if precio_ci  is not None else None
       um_v     = str_val(row.iloc[um_precio_ci]) if um_precio_ci is not None else ""


       rec = {k: (str_val(row.iloc[ci]) if ci < len(row) else "") for k, ci in col_map.items()}
       rec["concat1"] = sf + centro + destino + mat
       rec["concat2"] = sf + centro + mat
       rows.append(rec)


   df = pd.DataFrame(rows) if rows else pd.DataFrame()
   # Guardar metadatos de columnas para uso posterior
   df.attrs["col_map"]       = col_map
   df.attrs["sf_ci"]         = sf_ci
   df.attrs["nombre_sf_ci"]  = nombre_sf_ci
   df.attrs["mat_ci"]        = mat_ci
   df.attrs["centro_ci"]     = centro_ci
   df.attrs["destino_ci"]    = destino_ci
   df.attrs["precio_ci"]     = precio_ci
   df.attrs["um_precio_ci"]  = um_precio_ci
   df.attrs["col_offset"]    = col_offset
   df.attrs["hdr_idx"]       = hdr_idx
   print(f"   TRAOPE: {len(df)} registros")
   return df




def parse_materiales(path: Path) -> pd.DataFrame:
   """
   Parsea el catalogo de Materiales.


   - Solo conserva filas con Denom = 1000
   - Calcula PV = Contador / Denom  (no usa el PV que trae el archivo)


   Estructura esperada (con encabezados en fila 0):
     Material | TpMt | Denom. | UMA | Contador | PV | UMB | Texto breve ...
   """
   raw = read_excel_safe(path)


   # Localizar fila de encabezados
   hdr_idx = 0
   for i, row in raw.iterrows():
       vals = [str(v).strip().lower() for v in row if not (v is None or (not isinstance(v, str) and pd.isna(v)))]
       if "material" in vals:
           hdr_idx = i
           break


   hdr = raw.iloc[hdr_idx]
   # Usar lista ordenada para conservar primer ocurrencia en columnas duplicadas
   col_list = [(ci, str(v).strip()) for ci, v in enumerate(hdr)
               if not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip()]


   def find_col(*names) -> Optional[int]:
       """Devuelve el indice de la PRIMERA columna cuyo nombre contiene alguno de los terminos."""
       for n in names:
           for ci, k in col_list:
               if n.lower() in k.lower():
                   return ci
       return None


   mat_ci      = find_col("Material")
   tp_ci       = find_col("TpMt", "Tipo")
   denom_ci    = find_col("Denom")    # primera "Denom." = col 2 (valores 1000/1)
   uma_ci      = find_col("UMA")
   contador_ci = find_col("Contador")
   umb_ci      = find_col("UMB")
   nombre_ci   = find_col("Texto breve", "Denominacion", "Descripcion", "Nombre")


   rows = []
   for _, row in raw.iloc[hdr_idx + 1:].iterrows():
       if denom_ci is None:
           continue
       try:
           denom = float(row.iloc[denom_ci])
       except (TypeError, ValueError):
           continue
       if denom != 1000:
           continue


       mat = str_val(row.iloc[mat_ci]) if mat_ci is not None else ""
       if not mat or not re.match(r"^\d", mat):
           continue


       try:
           contador = float(row.iloc[contador_ci]) if contador_ci is not None else None
           pv = contador / 1000.0 if contador is not None else None
       except (TypeError, ValueError):
           pv = None


       rows.append({
           "material":  mat,
           "tp_mt":     str_val(row.iloc[tp_ci])     if tp_ci     is not None else "",
           "denom":     "1000",
           "uma":       str_val(row.iloc[uma_ci])    if uma_ci    is not None else "",
           "contador":  contador,
           "umb":       str_val(row.iloc[umb_ci])    if umb_ci    is not None else "",
           "pv":        pv,
           "nombre":    str_val(row.iloc[nombre_ci]) if nombre_ci is not None else "",
       })


   df = pd.DataFrame(rows)
   print(f"   Materiales: {len(df)} registros (solo Denom=1000)")
   return df




# ═══════════════════════════════════════════════════════════
#  CONSTRUCCION DE LOOKUPS EN PYTHON
# ═══════════════════════════════════════════════════════════


def build_traope_lookups(traope_df: pd.DataFrame) -> dict:
   """Construye diccionarios de busqueda a partir del TRAOPE."""
   sf_to_nombre:   Dict[str, str]   = {}
   dest_to_nombre: Dict[str, str]   = {}
   c1_to_precio:   Dict[str, float] = {}
   c2_to_precio:   Dict[str, float] = {}
   c1_to_um:       Dict[str, str]   = {}
   c2_to_um:       Dict[str, str]   = {}


   for _, row in traope_df.iterrows():
       sf      = str_val(row.get("Ship From", ""))
       nsf     = str_val(row.get("Nombre SF", ""))
       dest    = str_val(row.get("Destino", ""))
       ndest   = str_val(row.get("Nombre Destino", ""))
       precio  = row.get("Precio neto pedido")
       um      = str_val(row.get("UM Precio Pedido", ""))
       c1      = str_val(row.get("concat1", ""))
       c2      = str_val(row.get("concat2", ""))


       if sf and sf not in sf_to_nombre:
           sf_to_nombre[sf] = nsf
       if dest and dest not in dest_to_nombre:
           dest_to_nombre[dest] = ndest
       if c1 and c1 not in c1_to_precio:
           try:
               c1_to_precio[c1] = float(precio)
               c1_to_um[c1]     = um
           except (TypeError, ValueError):
               pass
       if c2 and c2 not in c2_to_precio:
           try:
               c2_to_precio[c2] = float(precio)
               c2_to_um[c2]     = um
           except (TypeError, ValueError):
               pass


   return {
       "sf_to_nombre":   sf_to_nombre,
       "dest_to_nombre": dest_to_nombre,
       "c1_to_precio":   c1_to_precio,
       "c2_to_precio":   c2_to_precio,
       "c1_to_um":       c1_to_um,
       "c2_to_um":       c2_to_um,
   }




def build_flete_lookups(flete_df: pd.DataFrame) -> dict:
   """Construye diccionarios de busqueda a partir del archivo de Fletes."""
   c1_to_cond:    Dict[str, str]   = {}
   c1_to_importe: Dict[str, float] = {}


   for _, row in flete_df.iterrows():
       sf   = str_val(row.get("sf", ""))
       ced  = str_val(row.get("cedis", ""))
       dest = str_val(row.get("destino", ""))
       mat  = str_val(row.get("material", ""))
       cond = str_val(row.get("cond_exp", ""))
       imp  = row.get("importe_flete")


       c1 = sf + ced + dest + mat
       if c1 and c1 not in c1_to_cond:
           c1_to_cond[c1] = cond
           try:
               c1_to_importe[c1] = float(imp)
           except (TypeError, ValueError):
               c1_to_importe[c1] = None


   return {"c1_to_cond": c1_to_cond, "c1_to_importe": c1_to_importe}




def build_mat_lookups(mat_df: pd.DataFrame) -> dict:
   """Construye diccionarios Material -> PV y Material -> Nombre."""
   pv: Dict[str, Optional[float]] = {}
   nombre: Dict[str, str] = {}
   for _, row in mat_df.iterrows():
       mat = str_val(row.get("material", ""))
       if mat:
           pv[mat]     = row.get("pv")
           nombre[mat] = str_val(row.get("nombre", ""))
   return {"pv": pv, "nombre": nombre}




# ═══════════════════════════════════════════════════════════
#  FORMULAS DE VALIDACION
# ═══════════════════════════════════════════════════════════


def validation_formula(r: int, cols: Dict[str, int] = COLS) -> str:
   """
   Logica de Validacion 1 (Excel):
     - Si UM Venta = TN  Y  UM Costo = TN  → MP + Flete
     - Si Cond Exp = 1  (UM costo en M3)   → PV * MP
     - Si Cond Exp = 4  (UM costo en M3)   → (MP + Flete) * PV
   """
   um_v = _ref("UM_VENTA",       r, cols)
   um_c = _ref("UM_COSTO",       r, cols)
   cond = _ref("COND_EXP",       r, cols)
   pv   = _ref("PV",             r, cols)
   mp   = _ref("IMPORTE_MP",     r, cols)
   fl   = _ref("IMPORTE_FLETE",  r, cols)
   # Nota: G&""="1" convierte tanto numero 1 como texto "1" a "1" para comparar correctamente
   return f'=IF(AND({um_v}="TN",{um_c}="TN"),{mp}+{fl},IF({cond}&""="1",{pv}*{mp},({mp}+{fl})*{pv}))'




def calc_validation1(um_venta, um_costo, cond_exp, pv, importe_mp, importe_flete) -> Optional[float]:
   """Calcula Validacion 1 en Python (para el archivo limpio)."""
   try:
       mp  = float(importe_mp)   if importe_mp   is not None else 0.0
       fl  = float(importe_flete) if importe_flete is not None else 0.0
       pv_ = float(pv)           if pv            is not None else 1.0
       if str(um_venta) == "TN" and str(um_costo) == "TN":
           return mp + fl
       elif str(cond_exp) == "1":
           return pv_ * mp
       else:
           return (mp + fl) * pv_
   except (TypeError, ValueError):
       return None




# ═══════════════════════════════════════════════════════════
#  ESCRITURA DE ESTILOS EN ENCABEZADOS
# ═══════════════════════════════════════════════════════════


def _write_section_headers(ws, cols: Dict[str, int] = COLS):
   """Fila 1: secciones VENTA y COSTO con color."""
   venta_start, venta_end = cols["IMPORTE_MP"],    cols["UM_VENTA"]
   costo_start, costo_end = cols["IMPORTE_COSTO"], cols["UM_COSTO"]
   ws.merge_cells(start_row=1, start_column=venta_start, end_row=1, end_column=venta_end)
   ws.merge_cells(start_row=1, start_column=costo_start, end_row=1, end_column=costo_end)
   for col_num, label, rgb in [(venta_start, "VENTA", CLR_VENTA_HDR), (costo_start, "COSTO", CLR_COSTO_HDR)]:
       c = ws.cell(row=1, column=col_num)
       c.value = label
       _apply(c,
              font=FONT_SEC,
              fill=_fill(rgb),
              alignment=ALIGN_CTR,
              border=BORDER)




def _write_col_headers(ws, cols: Dict[str, int] = COLS):
   """Fila 2: etiquetas de columna con color segun seccion."""
   for key, col_num in cols.items():
       cell = ws.cell(row=2, column=col_num)
       cell.value = COL_LABELS[key]
       if key in VENTA_KEYS:
           fill_rgb = CLR_VENTA_COL
       elif key in COSTO_KEYS:
           fill_rgb = CLR_COSTO_COL
       elif key in VAL_KEYS:
           fill_rgb = CLR_VAL_COL
       else:
           fill_rgb = CLR_HEADER
       _apply(cell, font=FONT_HDR, fill=_fill(fill_rgb), alignment=ALIGN_CTR, border=BORDER)




def _apply_col_dims(ws, cols: Dict[str, int] = COLS, freeze: str = "C3"):
   """Anchos de columna y alturas de filas."""
   for key, w in COL_WIDTHS.items():
       if key not in cols:
           continue
       ws.column_dimensions[_cl(key, cols)].width = w
   ws.row_dimensions[1].height = 20
   ws.row_dimensions[2].height = 32
   ws.freeze_panes = freeze




# ═══════════════════════════════════════════════════════════
#  SHEET PRINCIPAL CON FORMULAS (VLOOKUP)
# ═══════════════════════════════════════════════════════════


def build_main_sheet_formulas(ws, mp_df: pd.DataFrame, flete_lkp: dict, traope_lkp: dict, mat_lkp: dict):
   """
   Escribe el sheet principal con formulas VLOOKUP/IFERROR que apuntan
   a los sub-sheets: TRAOPE, Flete, Materiales.


   Referencias de columnas en sub-sheet TRAOPE (1-indexed, con Concat1 en A):
     F(6)  = Ship From      → VLOOKUP Nombre SF:  rango $F:$G, col 2
     V(22) = Precio neto    → VLOOKUP Importe Costo: rango $A:$V, col 22
     Y(25) = UM Precio Ped  → VLOOKUP UM Costo: rango $A:$Y, col 25
     AD(30)= Destino        → VLOOKUP Nombre Destino: rango $AD:$AE, col 2


   Referencias de columnas en sub-sheet Flete (1-indexed, con Concat1 en A):
     E(5)  = Cond. Exp.    → VLOOKUP: rango $A:$E, col 5
     J(10) = Importe Flete → VLOOKUP: rango $A:$J, col 10
   """
   _write_section_headers(ws)
   _write_col_headers(ws)


   def wv(col_key, row, value, fmt=None, align=None):
       """Escribe un valor con estilo."""
       cell = ws.cell(row=row, column=COLS[col_key])
       cell.value = value
       al = align or (ALIGN_RGT if fmt else ALIGN_LFT)
       _apply(cell, font=Font(size=10), alignment=al, border=BORDER)
       if fmt:
           cell.number_format = fmt


   def wf(col_key, row, formula, fmt=None, align=None):
       """Escribe una formula con estilo."""
       wv(col_key, row, formula, fmt, align)


   for idx, mp_row in mp_df.iterrows():
       r = idx + 3  # fila Excel


       sf      = mp_row["sf"]
       cedis   = mp_row["cedis"]
       destino = mp_row["destino"]
       mat     = mp_row["material"]


       # Columnas con valores directos de MP
       # NO_SF y NO_MATERIAL se escriben como enteros para que VLOOKUP coincida
       # con los numeros en TRAOPE!F y Materiales!A (que vienen del raw Excel como numeros)
       sf_num  = int(sf)  if sf.isdigit()  else sf
       mat_num = int(mat) if mat.isdigit() else mat


       wf("CONCAT1",         r, f'=CONCATENATE({_ref("NO_SF",r)},{_ref("CEDIS",r)},{_ref("DESTINO",r)},{_ref("NO_MATERIAL",r)})', align=ALIGN_LFT)
       wf("CONCAT2",         r, f'=CONCATENATE({_ref("NO_SF",r)},{_ref("CEDIS",r)},{_ref("NO_MATERIAL",r)})',                      align=ALIGN_LFT)
       wv("NO_SF",           r, sf_num,              align=ALIGN_CTR)
       wv("CEDIS",           r, cedis,               align=ALIGN_CTR)
       wv("DESTINO",         r, destino,             align=ALIGN_CTR)
       wv("NO_MATERIAL",     r, mat_num)
       wv("INICIO_VIG",      r, mp_row["inicio_vig"], align=ALIGN_CTR)
       wv("FIN_VIG",         r, mp_row["fin_vig"],    align=ALIGN_CTR)
       wv("IMPORTE_MP",      r, mp_row["importe_mp"], FMT_NUM2)
       wv("UM_VENTA",        r, mp_row["um"],         align=ALIGN_CTR)


       # BUSCARV: Nombre SF
       wf("NOMBRE_SF",      r, f'=IFERROR(VLOOKUP({_ref("NO_SF",r)},TRAOPE!$F:$G,2,0),"")')


       # BUSCARV: Condicion Expedicion (desde Flete Concat1)
       wf("COND_EXP",       r, f'=IFERROR(VLOOKUP({_ref("CONCAT1",r)},Flete!$A:$E,5,0),"")', align=ALIGN_CTR)


       # Nombre Destino: directo del MP si esta disponible, sino VLOOKUP desde TRAOPE
       nombre_dest_mp = mp_row.get("nombre_destino", "")
       if nombre_dest_mp:
           wv("NOMBRE_DESTINO", r, nombre_dest_mp)
       else:
           wf("NOMBRE_DESTINO", r, f'=IFERROR(VLOOKUP({_ref("DESTINO",r)},TRAOPE!$AD:$AE,2,0),"")')


       # BUSCARV: Nombre Material desde Materiales col H(8)
       wf("NOMBRE_MATERIAL", r, f'=IFERROR(VLOOKUP({_ref("NO_MATERIAL",r)},Materiales!$A:$H,8,0),"")')


       # BUSCARV: PV desde Materiales col G(7)
       wf("PV",             r, f'=IFERROR(VLOOKUP({_ref("NO_MATERIAL",r)},Materiales!$A:$G,7,0),"")', FMT_NUM3)


       # BUSCARV: Importe Flete
       wf("IMPORTE_FLETE",  r, f'=IFERROR(VLOOKUP({_ref("CONCAT1",r)},Flete!$A:$J,10,0),"")', FMT_NUM2)


       # BUSCARV: Importe Costo (IFERROR con Concat1 primero, luego Concat2)
       wf("IMPORTE_COSTO",  r,
           f'=IFERROR('
           f'VLOOKUP({_ref("CONCAT1",r)},TRAOPE!$A:$V,22,0),'
           f'VLOOKUP({_ref("CONCAT2",r)},TRAOPE!$A:$V,22,0))',
           FMT_NUM2)


       # BUSCARV: UM Costo (misma logica)
       wf("UM_COSTO",       r,
           f'=IFERROR('
           f'VLOOKUP({_ref("CONCAT1",r)},TRAOPE!$A:$Y,25,0),'
           f'VLOOKUP({_ref("CONCAT2",r)},TRAOPE!$A:$Y,25,0))',
           align=ALIGN_CTR)


       # Formulas de validacion
       wf("VALIDACION1", r, validation_formula(r), FMT_NUM2)
       wf("VALIDACION2", r, f'={_ref("IMPORTE_COSTO",r)}-{_ref("VALIDACION1",r)}', FMT_NUM2)


   _apply_col_dims(ws)




# ═══════════════════════════════════════════════════════════
#  SHEET PRINCIPAL CON VALORES (archivo entregable)
# ═══════════════════════════════════════════════════════════


def build_main_sheet_values(ws, mp_df: pd.DataFrame, flete_lkp: dict, traope_lkp: dict, mat_lkp: dict):
   """
   Escribe el sheet principal con valores calculados en Python.
   Las unicas formulas que permanecen son Validacion 1 y Validacion 2.
   Las celdas de Validacion 2 con diferencia > 1 se resaltan en amarillo.

   No incluye CONCAT1/CONCAT2 (llave 1 / llave 2): son solo llaves internas
   de cruce usadas para armar este archivo, no le sirven al usuario y aqui no
   hay formulas VLOOKUP que las necesiten (a diferencia del archivo "Con
   Formulas", donde si se conservan).
   """
   _write_section_headers(ws, COLS_CLEAN)
   _write_col_headers(ws, COLS_CLEAN)


   tf = traope_lkp
   ff = flete_lkp


   def wv(col_key, row, value, fmt=None, align=None, warn_fill=False):
       cell = ws.cell(row=row, column=COLS_CLEAN[col_key])
       cell.value = value
       al = align or (ALIGN_RGT if fmt else ALIGN_LFT)
       _apply(cell, font=Font(size=10), alignment=al, border=BORDER)
       if fmt:
           cell.number_format = fmt
       if warn_fill:
           cell.fill = _fill(CLR_WARN)


   for idx, mp_row in mp_df.iterrows():
       r = idx + 3


       sf      = mp_row["sf"]
       cedis   = mp_row["cedis"]
       destino = mp_row["destino"]
       mat     = mp_row["material"]


       concat1 = sf + cedis + destino + mat
       concat2 = sf + cedis + mat


       # Lookups
       nombre_sf     = tf["sf_to_nombre"].get(sf, "")
       nombre_dest   = mp_row.get("nombre_destino") or (tf["dest_to_nombre"].get(destino, "") if destino else "")
       cond_exp      = ff["c1_to_cond"].get(concat1, "")
       importe_flete = ff["c1_to_importe"].get(concat1)
       pv            = mat_lkp["pv"].get(mat)
       nombre_mat    = mat_lkp["nombre"].get(mat, "")
       importe_costo = tf["c1_to_precio"].get(concat1) or tf["c2_to_precio"].get(concat2)
       um_costo      = tf["c1_to_um"].get(concat1)    or tf["c2_to_um"].get(concat2) or ""
       importe_mp    = mp_row["importe_mp"]
       um_venta      = mp_row["um"]


       # Escribir valores (CONCAT1/CONCAT2 solo se usan arriba para los lookups,
       # no se escriben en este archivo -- ver docstring de la funcion)
       wv("NO_SF",           r, sf,                    align=ALIGN_CTR)
       wv("NOMBRE_SF",       r, nombre_sf)
       wv("CEDIS",           r, cedis,                 align=ALIGN_CTR)
       wv("DESTINO",         r, destino,               align=ALIGN_CTR)
       wv("COND_EXP",        r, cond_exp,              align=ALIGN_CTR)
       wv("NOMBRE_DESTINO",  r, nombre_dest)
       wv("NO_MATERIAL",     r, mat)
       wv("NOMBRE_MATERIAL", r, nombre_mat)
       wv("PV",              r, pv,                    FMT_NUM3)
       wv("INICIO_VIG",      r, mp_row["inicio_vig"],  align=ALIGN_CTR)
       wv("FIN_VIG",         r, mp_row["fin_vig"],     align=ALIGN_CTR)
       wv("IMPORTE_MP",      r, importe_mp,            FMT_NUM2)
       wv("IMPORTE_FLETE",   r, importe_flete,         FMT_NUM2)
       wv("UM_VENTA",        r, um_venta,              align=ALIGN_CTR)
       wv("IMPORTE_COSTO",   r, importe_costo,         FMT_NUM2)
       wv("UM_COSTO",        r, um_costo,              align=ALIGN_CTR)


       # Formulas de validacion (se conservan)
       wv("VALIDACION1", r, validation_formula(r, COLS_CLEAN), FMT_NUM2)
       wv("VALIDACION2", r,
          f'={_ref("IMPORTE_COSTO",r,COLS_CLEAN)}-{_ref("VALIDACION1",r,COLS_CLEAN)}',
          FMT_NUM2)


       # Calculo Python para highlight de diferencias > 1
       v1 = calc_validation1(um_venta, um_costo, cond_exp, pv, importe_mp, importe_flete)
       try:
           costo_f = float(importe_costo)
           if v1 is not None and abs(costo_f - v1) > 1:
               ws.cell(row=r, column=COLS_CLEAN["VALIDACION2"]).fill = _fill(CLR_WARN)
       except (TypeError, ValueError):
           pass


   _apply_col_dims(ws, COLS_CLEAN, freeze="A3")




# ═══════════════════════════════════════════════════════════
#  SUB-SHEETS: TRAOPE, FLETE, MATERIALES, MP
# ═══════════════════════════════════════════════════════════


def write_traope_sheet(ws, traope_path: Path, traope_df: pd.DataFrame):
   """
   Escribe el sub-sheet TRAOPE con Concat1 en columna A.


   Layout de columnas (1-indexed) en este sheet:
     A(1)   = Concat1 = ShipFrom+Centro+Destino+Material
     B(2)   = Doc. Compras
     ...
     F(6)   = Ship From         <- clave para VLOOKUP Nombre SF
     G(7)   = Nombre SF
     ...
     L(12)  = Material
     N(14)  = Centro
     ...
     V(22)  = Precio neto pedido <- clave para VLOOKUP Importe Costo
     Y(25)  = UM Precio Pedido   <- clave para VLOOKUP UM Costo
     AD(30) = Destino            <- clave para VLOOKUP Nombre Destino
     AE(31) = Nombre Destino
   """
   raw = read_excel_safe(traope_path)
   hdr_idx    = traope_df.attrs.get("hdr_idx", 0)
   col_offset = traope_df.attrs.get("col_offset", 0)


   hdr_series = raw.iloc[hdr_idx]
   ordered_headers = [(ci, str(v).strip()) for ci, v in enumerate(hdr_series)
                      if ci >= col_offset and not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip()]


   # Determinar indices de columnas clave en raw
   def find_raw_ci(*names):
       for n in names:
           for ci, h in ordered_headers:
               if n.lower() in h.lower():
                   return ci
       return None


   sf_ci     = find_raw_ci("Ship From")
   mat_ci    = find_raw_ci("Material")
   centro_ci = find_raw_ci("Centro")
   dest_ci   = find_raw_ci("Destino")


   # Inicio real de datos
   data_start = hdr_idx + 1
   for i in range(data_start, min(data_start + 5, len(raw))):
       non_null = sum(1 for v in raw.iloc[i] if not (v is None or (not isinstance(v, str) and pd.isna(v))) and str(v).strip())
       if non_null >= 5:
           data_start = i
           break


   # Encabezado fila 1
   hdr_cell = ws.cell(row=1, column=1)
   hdr_cell.value = "Concat1"
   _apply(hdr_cell, font=FONT_HDR, fill=_fill(CLR_HEADER), alignment=ALIGN_CTR, border=BORDER)


   for out_col, (_, hdr_name) in enumerate(ordered_headers, start=2):
       cell = ws.cell(row=1, column=out_col)
       cell.value = hdr_name
       _apply(cell, font=FONT_HDR, fill=_fill(CLR_HEADER), alignment=ALIGN_CTR, border=BORDER)


   # Datos
   out_row = 2
   for _, row in raw.iloc[data_start:].iterrows():
       sf    = str_val(row.iloc[sf_ci])    if sf_ci    is not None else ""
       mat   = str_val(row.iloc[mat_ci])   if mat_ci   is not None else ""
       ctr   = str_val(row.iloc[centro_ci]) if centro_ci is not None else ""
       dest  = str_val(row.iloc[dest_ci])  if dest_ci  is not None else ""


       if not sf or not re.match(r"^\d", sf):
           continue


       concat1 = sf + ctr + dest + mat
       ws.cell(row=out_row, column=1).value = concat1


       for out_col, (src_ci, _) in enumerate(ordered_headers, start=2):
           v = row.iloc[src_ci] if src_ci < len(row) else None
           if isinstance(v, (pd.Timestamp, datetime)):
               v = v.strftime("%d.%m.%Y")
           elif not isinstance(v, str) and pd.isna(v):
               v = None
           ws.cell(row=out_row, column=out_col).value = v


       out_row += 1


   # Anchos
   ws.column_dimensions["A"].width = 35
   for i in range(2, len(ordered_headers) + 2):
       ws.column_dimensions[get_column_letter(i)].width = 14
   ws.freeze_panes = "B2"




def write_flete_sheet(ws, flete_path: Path, flete_df: pd.DataFrame):
   """
   Escribe el sub-sheet Flete con Concat1 en columna A.


   Layout (1-indexed):
     A(1)  = Concat1 = SF+Cedis+Destino+Material
     B(2)  = SF
     C(3)  = Cedis
     D(4)  = Destino
     E(5)  = Cond. Exp.   <- VLOOKUP col 5 desde $A:$E
     F(6)  = Material
     G(7)  = (blanco)
     H(8)  = Nombre Material
     I(9)  = (blanco)
     J(10) = Importe Flete <- VLOOKUP col 10 desde $A:$J
     ...
   """
   flete_col_headers = [
       "Concat1", "No. SF", "Cedis", "Destino", "Cond. Exp.", "Material",
       "", "Nombre Material", "", "Importe Flete", "Moneda",
       "Cantidad", "UM", "C", "", "Inicio Vigencia", "Fin Vigencia"
   ]


   for col, hdr in enumerate(flete_col_headers, start=1):
       cell = ws.cell(row=1, column=col)
       cell.value = hdr
       if hdr:
           _apply(cell, font=FONT_HDR, fill=_fill(CLR_HEADER), alignment=ALIGN_CTR, border=BORDER)


   out_row = 2


   if flete_df.attrs.get("is_vertical", False):
       # Nuevo formato: escribir desde los datos ya parseados
       for _, row in flete_df.iterrows():
           sf = str_val(row.get("sf", ""))
           if not sf:
               continue
           cedis   = str_val(row.get("cedis", ""))
           destino = str_val(row.get("destino", ""))
           mat     = str_val(row.get("material", ""))
           concat1 = sf + cedis + destino + mat


           ws.cell(row=out_row, column=1).value  = concat1
           ws.cell(row=out_row, column=2).value  = sf
           ws.cell(row=out_row, column=3).value  = cedis
           ws.cell(row=out_row, column=4).value  = destino
           ws.cell(row=out_row, column=5).value  = str_val(row.get("cond_exp", ""))
           ws.cell(row=out_row, column=6).value  = mat
           ws.cell(row=out_row, column=7).value  = None   # blanco
           ws.cell(row=out_row, column=8).value  = str_val(row.get("nombre_material", ""))
           ws.cell(row=out_row, column=9).value  = None   # blanco
           ws.cell(row=out_row, column=10).value = row.get("importe_flete")
           ws.cell(row=out_row, column=11).value = str_val(row.get("moneda", ""))
           ws.cell(row=out_row, column=12).value = row.get("cantidad")
           ws.cell(row=out_row, column=13).value = str_val(row.get("um", ""))
           ws.cell(row=out_row, column=14).value = str_val(row.get("sin_nombre", ""))
           ws.cell(row=out_row, column=15).value = None   # blanco
           ws.cell(row=out_row, column=16).value = str_val(row.get("inicio_vig", ""))
           ws.cell(row=out_row, column=17).value = str_val(row.get("fin_vig", ""))
           out_row += 1
   else:
       # Formato antiguo: copiar columnas del archivo raw directamente
       raw = read_excel_safe(flete_path)
       start = find_data_start_row(raw)


       for _, row in raw.iloc[start:].iterrows():
           sf = str_val(row.iloc[0])
           if not sf or not re.match(r"^\d", sf):
               continue


           cedis   = str_val(row.iloc[1])
           destino = str_val(row.iloc[2])
           mat     = str_val(row.iloc[4])
           concat1 = sf + cedis + destino + mat


           ws.cell(row=out_row, column=1).value = concat1


           # Copiar columnas raw 0..15 → columnas B..Q (out_col 2..17)
           for out_col, src_ci in enumerate(range(16), start=2):
               if src_ci < len(row):
                   v = row.iloc[src_ci]
                   if isinstance(v, (pd.Timestamp, datetime)):
                       v = v.strftime("%d.%m.%Y")
                   elif not isinstance(v, str) and pd.isna(v):
                       v = None
                   ws.cell(row=out_row, column=out_col).value = v
           out_row += 1


   ws.column_dimensions["A"].width = 35
   for i in range(2, 18):
       ws.column_dimensions[get_column_letter(i)].width = 13
   ws.freeze_panes = "B2"




def write_materiales_sheet(ws, mat_df: pd.DataFrame):
   """
   Escribe el sub-sheet Materiales limpio.


   Layout (1-indexed):
     A(1) = Material  <- clave VLOOKUP para PV
     B(2) = TpMt
     C(3) = Denom
     D(4) = UMA
     E(5) = Contador
     F(6) = UMB
     G(7) = PV = Contador/1000   <- VLOOKUP col 7 desde $A:$G
   """
   headers = ["Material", "TpMt", "Denom", "UMA", "Contador", "UMB", "PV", "Nombre Material"]
   widths  = [14, 8, 8, 6, 10, 6, 8, 45]


   for col, (hdr, w) in enumerate(zip(headers, widths), start=1):
       cell = ws.cell(row=1, column=col)
       cell.value = hdr
       _apply(cell, font=FONT_HDR, fill=_fill(CLR_HEADER), alignment=ALIGN_CTR, border=BORDER)
       ws.column_dimensions[get_column_letter(col)].width = w


   for out_row, (_, mrow) in enumerate(mat_df.iterrows(), start=2):
       pv = mrow.get("pv")
       mat_val = mrow.get("material", "")
       # Material como entero para que VLOOKUP(numero, Materiales!A, ...) coincida
       try:
           mat_val = int(mat_val)
       except (ValueError, TypeError):
           pass
       vals = [
           mat_val,
           mrow.get("tp_mt",    ""),
           mrow.get("denom",  "1000"),
           mrow.get("uma",      ""),
           mrow.get("contador"),
           mrow.get("umb",      ""),
           round(pv, 4) if pv is not None else "",
           mrow.get("nombre",   ""),
       ]
       for col, v in enumerate(vals, start=1):
           ws.cell(row=out_row, column=col).value = v


   ws.freeze_panes = "B2"




def write_mp_sheet(ws, mp_df: pd.DataFrame):
   """Escribe el sub-sheet MP limpio con encabezados y formato."""
   headers = [
       "No. SF", "Cedis", "Destino", "No. Material", "Nombre Material",
       "Importe MP", "Moneda", "Cantidad", "UM", "C",
       "Inicio Vigencia", "Fin Vigencia"
   ]
   widths = [14, 8, 10, 14, 40, 14, 8, 8, 6, 4, 14, 14]


   for col, (hdr, w) in enumerate(zip(headers, widths), start=1):
       cell = ws.cell(row=1, column=col)
       cell.value = hdr
       _apply(cell, font=FONT_HDR, fill=_fill(CLR_HEADER), alignment=ALIGN_CTR, border=BORDER)
       ws.column_dimensions[get_column_letter(col)].width = w


   for out_row, (_, row) in enumerate(mp_df.iterrows(), start=2):
       vals = [
           row["sf"],     row["cedis"],    row["destino"],
           row["material"], row["nombre_material"],
           row["importe_mp"], row["moneda"], row["cantidad"],
           row["um"],     row["sin_nombre"],
           row["inicio_vig"], row["fin_vig"]
       ]
       for col, v in enumerate(vals, start=1):
           ws.cell(row=out_row, column=col).value = v
           if col == 6 and v is not None:  # Importe MP
               ws.cell(row=out_row, column=col).number_format = FMT_NUM2


   ws.freeze_panes = "B2"




# ═══════════════════════════════════════════════════════════
#  CONSTRUCCION COMPLETA DEL LIBRO
# ═══════════════════════════════════════════════════════════


def build_workbook(
   mp_df:      pd.DataFrame,
   flete_df:   pd.DataFrame,
   traope_df:  pd.DataFrame,
   mat_df:     pd.DataFrame,
   cedis:      str,
   raw_files:  Dict[str, Path],
   formula_mode: bool = True,
) -> openpyxl.Workbook:
   """Construye el workbook completo con todos los sheets."""


   # Lookups Python
   traope_lkp = build_traope_lookups(traope_df)
   flete_lkp  = build_flete_lookups(flete_df)
   mat_lkp    = build_mat_lookups(mat_df)


   wb = openpyxl.Workbook()


   # Sheet principal
   ws_main = wb.active
   ws_main.title = f"Base Cedis {cedis}"


   if formula_mode:
       build_main_sheet_formulas(ws_main, mp_df, flete_lkp, traope_lkp, mat_lkp)
   else:
       build_main_sheet_values(ws_main, mp_df, flete_lkp, traope_lkp, mat_lkp)


   # Sub-sheets de soporte
   ws_mp  = wb.create_sheet("MP")
   write_mp_sheet(ws_mp, mp_df)


   ws_fl  = wb.create_sheet("Flete")
   write_flete_sheet(ws_fl, raw_files["flete"], flete_df)


   ws_mat = wb.create_sheet("Materiales")
   write_materiales_sheet(ws_mat, mat_df)


   ws_tr  = wb.create_sheet("TRAOPE")
   write_traope_sheet(ws_tr, raw_files["traope"], traope_df)


   return wb




# ═══════════════════════════════════════════════════════════
#  CLI PRINCIPAL
# ═══════════════════════════════════════════════════════════


def main():
   parser = argparse.ArgumentParser(
       description="CEMEX - Genera Base Cedis DW## a partir de 4 archivos de entrada",
       formatter_class=argparse.RawDescriptionHelpFormatter,
       epilog="""
Ejemplos:
 python3 generar_matriz.py /ruta/a/carpeta
 python3 generar_matriz.py --mp "DW88 MP.xlsx" --flete "DW88 Fletes.xlsx" --traope "DW88 TRAOPE.xlsx" --materiales "DW88 MATERIALES.xlsx"
 python3 generar_matriz.py /ruta/carpeta --output /ruta/salida --cedis DW88
       """
   )
   parser.add_argument("carpeta", nargs="?",      help="Carpeta con los 4 archivos de entrada")
   parser.add_argument("--mp",          metavar="FILE", help="Archivo MP / precios de venta / VK13")
   parser.add_argument("--flete",       metavar="FILE", help="Archivo de Fletes")
   parser.add_argument("--traope",      metavar="FILE", help="Archivo TRAOPE")
   parser.add_argument("--materiales",  metavar="FILE", help="Catalogo de Materiales")
   parser.add_argument("--output", "-o",metavar="DIR",  help="Carpeta de salida (default: carpeta de entrada)")
   parser.add_argument("--cedis",       metavar="CODE", help="Codigo CEDIS, ej: DW88 (auto-detectado si se omite)")
   args = parser.parse_args()


   # ── Detectar archivos ──
   raw_files: Dict[str, Path] = {}


   if args.carpeta:
       folder = Path(args.carpeta)
       if not folder.is_dir():
           parser.error(f"La carpeta no existe: {folder}")
       raw_files = find_files_in_folder(folder)


   # Argumentos explícitos sobrescriben la deteccion automatica
   for ftype in ("mp", "flete", "traope", "materiales"):
       val = getattr(args, ftype)
       if val:
           raw_files[ftype] = Path(val)


   missing = [t for t in ("mp", "flete", "traope", "materiales") if t not in raw_files]
   if missing:
       print(f"\nError: No se encontraron archivos para: {', '.join(missing)}", file=sys.stderr)
       if args.carpeta:
           folder = Path(args.carpeta)
           print("\nArchivos en la carpeta:", file=sys.stderr)
           for p in sorted(folder.iterdir()):
               if p.suffix.lower() in EXCEL_EXTS and not p.name.startswith("~"):
                   ftype = identify_file(p)
                   print(f"  {p.name:45s} → {ftype or '(no reconocido)'}", file=sys.stderr)
       print("\nEspecifica los archivos manualmente con --mp, --flete, --traope, --materiales", file=sys.stderr)
       sys.exit(1)


   print("Archivos detectados:")
   for ftype in ("mp", "flete", "traope", "materiales"):
       print(f"  {ftype:12s} → {raw_files[ftype].name}")


   cedis = args.cedis or extract_cedis(raw_files)
   print(f"\nCEDIS detectado: {cedis}")


   # ── Parsear ──
   print("\nParsando archivos...")
   mp_df      = parse_mp(raw_files["mp"])
   flete_df   = parse_flete(raw_files["flete"])
   traope_df  = parse_traope(raw_files["traope"])
   mat_df     = parse_materiales(raw_files["materiales"])


   # ── Carpeta de salida ──
   out_dir = Path(args.output) if args.output else raw_files["mp"].parent
   out_dir.mkdir(parents=True, exist_ok=True)


   # ── Generar libros ──
   print("\nGenerando archivos Excel...")


   out_formulas = out_dir / f"Base Cedis {cedis} (Con Formulas).xlsx"
   print(f"  Con formulas: {out_formulas.name} ...", end="", flush=True)
   wb_f = build_workbook(mp_df, flete_df, traope_df, mat_df, cedis, raw_files, formula_mode=True)
   wb_f.save(out_formulas)
   print(" OK")


   out_clean = out_dir / f"Base Cedis {cedis}.xlsx"
   print(f"  Limpio:        {out_clean.name} ...", end="", flush=True)
   wb_c = build_workbook(mp_df, flete_df, traope_df, mat_df, cedis, raw_files, formula_mode=False)
   wb_c.save(out_clean)
   print(" OK")


   print(f"\nArchivos guardados en: {out_dir}")
   print(f"  {out_formulas.name}")
   print(f"  {out_clean.name}")




if __name__ == "__main__":
   main()




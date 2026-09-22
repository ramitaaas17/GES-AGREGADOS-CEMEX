# -*- coding: utf-8 -*-
"""
Base de datos propia de rutas, construida a partir de la hoja BASE 2026 (tabla "Solicitudes":
historial de altas, cambios y bajas de rutas) del libro de extracción.

Sirve para completar datos que SAP/Snowflake no trae en algunas rutas:
    - Cond. Expedición   (en BASE 2026 se llama MODALIDAD: ´01 -> 1, ´02 -> 2, ´04 -> 4)
    - Nombre SF, Nombre del Centro, Nombre del Destino (planta / obra), Descripción del material
    - PV (factor peso/volumen) cuando el material no está en PESO_VOL

La base vive en la hoja BD_Completa del libro maestro (CEMEX_MATRIZ_INTEGRADO.xlsm) y es EDITABLE: se puede
corregir una celda o agregar filas (con Centro, Ship From, Destino y Material; una fecha más nueva gana). El motor
la lee de ahí (bd_desde_hoja). agregar_bd_a_master.py la (re)construye desde BASE 2026 + fuentes de apoyo.

Reglas de búsqueda (BuscadorBD):
    1. SAP manda: la base solo llena huecos y siempre se indica cuáles campos se completaron.
    2. Se usa el registro MÁS RECIENTE de cada ruta (SF + Centro + Destino + Material), ordenando por
       Fecha de Procesamiento (o Fecha de Solicitud si falta) y, en empate, por fila en la hoja.
    3. Movimientos BAJA / CANCELADA no se usan como fuente de datos de la ruta.
    4. Si la ruta exacta no está, se busca por SF + Centro + Material (sin destino), pero SOLO si todos
       los registros encontrados coinciden en ese valor (si hay ambigüedad no se inventa nada).
    5. Nombres (SF, Centro, Destino, Material): el más FRECUENTE registrado para ese código (así una errata
       suelta, p. ej. "LARARO" por "LAZARO", no gana); en caso de empate, el más reciente.
    6. Fuentes de apoyo (carpeta fuentes_apoyo, p. ej. "PRUEBA 1.xlsx": hojas "Formato Alta Precio" y
       "Materiales,canteras") se suman a la base con prioridad menor. Sus registros no traen SF, así que se
       buscan por Centro + Destino + Material (y Centro + Material para el PV), también solo si el valor es único.

Fechas: BASE 2026 mezcla texto dd.mm.aaaa con fechas reales; se interpretan SIEMPRE como día/mes/año
(pandas por defecto las leería como mes/día y mostraría "fechas futuras" falsas).
"""
import glob
import logging
import os
import re
import unicodedata
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HOJA_BASE = 'BASE 2026'
FILA_ENCABEZADO = 3          # índice para pandas: el encabezado está en la fila 4 de Excel
TIPOS_BAJA = {'BAJA', 'CANCELADA'}

# Encabezado en BASE 2026 (normalizado: sin acentos, mayúsculas) -> nombre limpio en la base propia
COLUMNAS_ORIGEN = {
    'RUTA': 'Ruta',
    'ORG VTAS': 'Org. Ventas',
    'CEDIS': 'Centro',
    'NOMBRE CEDIS': 'Nombre Cedis',
    'SF': 'Ship From',
    'NOMBRE SF': 'Nombre SF',
    'NO PLANTA / OBRA': 'Destino',
    'NOMBRE PLANTA / OBRA': 'Nombre Destino',
    'ARTICULO': 'Material',
    'DESC/ ARTICULO': 'Desc. Material',
    'DESC/ARTICULO': 'Desc. Material',
    'PV': 'PV',
    'MODALIDAD': 'Modalidad (original)',
    'UM $ VTA': 'UM Venta',
    'TIPO': 'Tipo Movimiento',
    'FECHA DE PROCESAMIENTO': 'Fecha Procesamiento',
    'TIPO CLIENTE': 'Tipo Cliente',
    'FEECHA SOLICITUD': 'Fecha Solicitud',      # así viene escrito en la hoja
    'FECHA SOLICITUD': 'Fecha Solicitud',
    'ESTADO DE SOLICITUD': 'Estado Solicitud',
    'OBSERVACION': 'Observación',
}
COLUMNAS_LLAVE = ('Centro', 'Ship From', 'Destino', 'Material')
COLUMNAS_TEXTO = ('Nombre Cedis', 'Nombre SF', 'Nombre Destino', 'Desc. Material', 'Ruta', 'Observación')
COLUMNAS_BD = [
    'Fuente', 'Fila Excel', 'Llave', 'Llave sin destino', 'Llave sin SF', 'Llave centro-material', 'Zona', 'Org. Ventas', 'Centro', 'Nombre Cedis', 'Ship From', 'Nombre SF',
    'Destino', 'Nombre Destino', 'Material', 'Desc. Material', 'PV', 'Modalidad (original)', 'Cond. Expedición',
    'UM Venta', 'Tipo Movimiento', 'Tipo Cliente', 'Estado Solicitud', 'Fecha Procesamiento', 'Fecha Solicitud',
    'Es baja', 'Registro usado en búsqueda', 'Ruta', 'Observación',
]


# ---------------------------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------------------------
def _norm(texto):
    t = unicodedata.normalize('NFKD', str(texto))
    t = ''.join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', t).strip().upper()


COLUMNAS_ORIGEN_NORM = {_norm(k): v for k, v in COLUMNAS_ORIGEN.items()}


def _vacio(x):
    return x is None or (isinstance(x, float) and pd.isna(x)) or str(x).strip().lower() in ('', 'nan', 'nat', 'none')


def _cod(x):
    """'10000617.0' -> '10000617'; ' D838 ' -> 'D838'; vacío -> ''."""
    if _vacio(x):
        return ''
    s = str(x).strip()
    if re.fullmatch(r'\d+\.0+', s):
        s = s.split('.')[0]
    return s


def _texto(x):
    return '' if _vacio(x) else re.sub(r'\s+', ' ', str(x)).strip()


def _fecha(x):
    """Interpreta siempre día/mes/año. Acepta '2026-01-06 00:00:00', '10.09.2026', '10/09/2026' y seriales de Excel."""
    if _vacio(x):
        return pd.NaT
    s = str(x).strip()
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        return pd.to_datetime(f'{m[1]}-{int(m[2]):02d}-{int(m[3]):02d}', errors='coerce')
    m = re.match(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{4})', s)
    if m:
        return pd.to_datetime(f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}', errors='coerce')
    try:
        v = float(s)
        if 30000 < v < 80000:
            return pd.Timestamp('1899-12-30') + pd.Timedelta(days=v)
    except ValueError:
        pass
    return pd.NaT


def _modalidad(x):
    """'´04' -> '4'; '´01' -> '1'. Si trae otro código se conserva (para que aparezca en Calidad)."""
    if _vacio(x):
        return ''
    d = re.sub(r'\D', '', str(x)).lstrip('0')
    return d


# ---------------------------------------------------------------------------------------------
# Construcción de la base
# ---------------------------------------------------------------------------------------------
def _finalizar(bd, respetar_cond=False):
    """Limpieza y columnas derivadas comunes a BASE 2026 y a las fuentes de apoyo."""
    for c in COLUMNAS_BD:
        if c not in bd.columns:
            bd[c] = ''
    for c in COLUMNAS_LLAVE + ('Org. Ventas',):
        bd[c] = bd[c].map(_cod)
    for c in COLUMNAS_TEXTO + ('Zona',):
        bd[c] = bd[c].map(_texto)
    for c in ('Tipo Movimiento', 'Tipo Cliente', 'Estado Solicitud', 'UM Venta'):
        bd[c] = bd[c].map(_texto).str.upper()
    bd = bd[(bd['Centro'] != '') & (bd['Material'] != '')].copy()

    cond_orig = bd['Modalidad (original)'].map(_modalidad)
    if respetar_cond:      # hoja BD_Completa: lo escrito en "Cond. Expedición" (posible corrección manual) manda
        cond_hoja = bd['Cond. Expedición'].map(_modalidad)
        bd['Cond. Expedición'] = cond_hoja.where(cond_hoja != '', cond_orig)
    else:
        bd['Cond. Expedición'] = cond_orig
    bd['Fila Excel'] = pd.to_numeric(bd['Fila Excel'], errors='coerce').fillna(0).astype(int)
    bd['PV'] = pd.to_numeric(bd['PV'], errors='coerce').round(3)
    bd['Fecha Procesamiento'] = bd['Fecha Procesamiento'].map(_fecha)
    bd['Fecha Solicitud'] = bd['Fecha Solicitud'].map(_fecha)
    bd['Llave'] = bd['Ship From'] + '|' + bd['Centro'] + '|' + bd['Destino'] + '|' + bd['Material']
    bd['Llave sin destino'] = bd['Ship From'] + '|' + bd['Centro'] + '|' + bd['Material']
    bd['Llave sin SF'] = bd['Centro'] + '|' + bd['Destino'] + '|' + bd['Material']
    bd['Llave centro-material'] = bd['Centro'] + '|' + bd['Material']
    bd['Es baja'] = bd['Tipo Movimiento'].isin(TIPOS_BAJA)
    bd['_orden'] = bd['Fecha Procesamiento'].fillna(bd['Fecha Solicitud']).fillna(pd.Timestamp('1900-01-01'))
    return bd


def _base_2026(df_raw):
    """Hoja BASE 2026 (encabezado en fila 4) -> filas con nombres limpios, o None."""
    if df_raw is None or len(df_raw) == 0:
        return None
    mapa = {}
    for c in df_raw.columns:
        destino = COLUMNAS_ORIGEN_NORM.get(_norm(c))
        if destino and destino not in mapa.values():
            mapa[c] = destino
    faltan = [c for c in COLUMNAS_LLAVE if c not in mapa.values()]
    if faltan:
        logging.warning(f"BASE 2026: faltan columnas clave {faltan}; se ignora esa hoja.")
        return None
    bd = df_raw[list(mapa)].rename(columns=mapa).copy()
    bd['Fila Excel'] = [i + 5 for i in range(len(bd))]     # encabezado en la fila 4 -> primer dato en la 5
    bd['Fuente'] = 'BASE 2026'
    return bd


def _col(df, *prefijos):
    """Primera columna cuyo nombre normalizado empieza con alguno de los prefijos."""
    for c in df.columns:
        n = _norm(c)
        if any(n.startswith(p) for p in prefijos):
            return c
    return None


def _hoja(xls, nombre):
    objetivo = _norm(nombre)
    return next((h for h in xls.sheet_names if _norm(h) == objetivo), None)


def _leer_alta_precio(xls, hoja, archivo):
    df = pd.read_excel(xls, sheet_name=hoja, dtype=str)
    c = {
        'Centro': _col(df, 'CENTRO'), 'Nombre Cedis': _col(df, 'NOMBRE PLANTA'), 'Zona': _col(df, 'NUEVA ZONA'),
        'Destino': _col(df, 'DEST.MERC'), 'Nombre Destino': _col(df, 'NOMBRE DEL DEST'),
        'Material': _col(df, 'MATERIAL'), 'Desc. Material': _col(df, 'NOMBRE DE MATERIAL'), 'PV': _col(df, 'PV'),
        'UM Venta': _col(df, 'UM'), 'Modalidad (original)': _col(df, 'MODALIDAD'),
        'Fecha Procesamiento': _col(df, 'VALIDO DE'),
    }
    if not c['Centro'] or not c['Material']:
        logging.warning(f"{archivo} / {hoja}: no se reconocieron las columnas Centro / Material; se ignora.")
        return None
    out = pd.DataFrame({k: (df[v] if v else '') for k, v in c.items()})
    out['Fila Excel'] = [i + 2 for i in range(len(out))]
    out['Tipo Movimiento'] = 'ALTA'
    out['Fuente'] = f'{archivo} / {hoja}'
    return out


def _leer_materiales_canteras(xls, hoja, archivo):
    df = pd.read_excel(xls, sheet_name=hoja, dtype=str)
    c = {'Centro': _col(df, 'NO. CANTERA', 'CENTRO'), 'Nombre Cedis': _col(df, 'NOMBRE CANTERA'),
         'Material': _col(df, 'MATERIAL'), 'Desc. Material': _col(df, 'DENOMINACION'), 'PV': _col(df, 'PV')}
    if not c['Centro'] or not c['Material']:
        logging.warning(f"{archivo} / {hoja}: no se reconocieron las columnas Cantera / Material; se ignora.")
        return None
    out = pd.DataFrame({k: (df[v] if v else '') for k, v in c.items()})
    out['Fila Excel'] = [i + 2 for i in range(len(out))]
    out['Tipo Movimiento'] = 'CATALOGO'
    out['Fuente'] = f'{archivo} / {hoja}'
    return out


def cargar_apoyo(ruta):
    """Lee las fuentes de apoyo. ruta: archivo .xlsx o carpeta (se leen los .xlsx que traigan hojas conocidas).
    Devuelve un DataFrame con nombres de la base propia, o None."""
    if not ruta or not os.path.exists(ruta):
        return None
    archivos = [ruta] if os.path.isfile(ruta) else sorted(glob.glob(os.path.join(ruta, '*.xlsx')))
    partes = []
    for f in archivos:
        nombre = os.path.basename(f)
        if nombre.startswith('~$'):
            continue
        try:
            xls = pd.ExcelFile(f)
        except Exception as e:
            logging.warning(f"Fuente de apoyo ilegible ({nombre}): {e}")
            continue
        h1, h2 = _hoja(xls, 'Formato Alta Precio'), _hoja(xls, 'Materiales,canteras')
        if not h1 and not h2:
            logging.info(f"Fuente de apoyo omitida (sin hojas conocidas): {nombre}")
            continue
        for leer, hoja in ((_leer_alta_precio, h1), (_leer_materiales_canteras, h2)):
            if hoja:
                parte = leer(xls, hoja, nombre)
                if parte is not None:
                    partes.append(parte)
    return pd.concat(partes, ignore_index=True) if partes else None


def construir_bd(df_raw, apoyo=None):
    """Base propia = BASE 2026 (df_raw, encabezado en fila 4) + fuentes de apoyo (opcional). None si no hay nada."""
    partes = [p for p in (_base_2026(df_raw), apoyo) if p is not None and len(p)]
    if not partes:
        logging.warning("Sin BASE 2026 ni fuentes de apoyo: no se construye la base propia.")
        return None
    bd = _finalizar(pd.concat(partes, ignore_index=True))
    if bd.empty:
        return None

    return _marcar_usados(bd)


def _marcar_usados(bd):
    """Marca el registro más reciente no-baja de cada ruta (el que usa la búsqueda exacta)."""
    vigentes = bd[~bd['Es baja']].sort_values(['_orden', 'Fila Excel'])
    usados = set(vigentes.drop_duplicates('Llave', keep='last').index)
    bd['Registro usado en búsqueda'] = bd.index.isin(usados)
    return bd.reset_index(drop=True)


def bd_desde_hoja(df_hoja):
    """Hoja BD_Completa del libro (encabezado en fila 1) -> base lista para buscar; None si la hoja no sirve.
    Se respeta lo escrito en "Cond. Expedición"; llaves, bajas y "registro usado" se recalculan."""
    if df_hoja is None or len(df_hoja) == 0:
        return None
    bd = df_hoja.copy()
    bd.columns = [str(c).strip() for c in bd.columns]
    faltan = [c for c in COLUMNAS_LLAVE if c not in bd.columns]
    if faltan:
        logging.warning(f"Hoja BD_Completa: faltan columnas {faltan}; se ignora.")
        return None
    bd = _finalizar(bd, respetar_cond=True)
    if bd.empty:
        return None
    bd['Fuente'] = bd['Fuente'].map(_texto).replace('', 'Manual (BD_Completa)')
    return _marcar_usados(bd)


class BuscadorBD:
    """Índices en memoria para completar campos de la matriz (ver reglas en la cabecera del módulo)."""

    def __init__(self, bd):
        self.bd = bd
        orden = ['_orden', 'Fila Excel']
        todos = bd.sort_values(orden)
        v = todos[~todos['Es baja']]

        self._l1 = {r['Llave']: r for r in v.drop_duplicates('Llave', keep='last').to_dict('records')}

        def unico(col):
            w = v[v[col].notna() & (v[col] != '')]
            g = w.groupby('Llave sin destino')[col].agg(lambda s: set(s))
            return {k: next(iter(s)) for k, s in g.items() if len(s) == 1}

        self._cond3 = unico('Cond. Expedición')
        self._pv3 = unico('PV')

        def unico_por(col, llave):
            w = v[v[col].notna() & (v[col] != '') & (v[llave] != '')]
            g = w.groupby(llave)[col].agg(lambda s: set(s))
            return {k: next(iter(s)) for k, s in g.items() if len(s) == 1}

        # Sin SF (las fuentes de apoyo no lo traen): Centro + Destino + Material  /  Centro + Material
        self._cond_csd = unico_por('Cond. Expedición', 'Llave sin SF')
        self._pv_csd = unico_por('PV', 'Llave sin SF')
        self._pv_cm = unico_por('PV', 'Llave centro-material')

        def ultimo(col_codigo, col_nombre):
            """Nombre más frecuente registrado para el código (una errata suelta no gana); en empate, el más reciente."""
            w = todos[(todos[col_codigo] != '') & (todos[col_nombre] != '')]
            cnt = (w.groupby([col_codigo, col_nombre])
                    .agg(n=('Fila Excel', 'size'), ult=('_orden', 'max')).reset_index()
                    .sort_values(['n', 'ult']).drop_duplicates(col_codigo, keep='last'))
            return dict(zip(cnt[col_codigo], cnt[col_nombre]))

        self._nom_sf = ultimo('Ship From', 'Nombre SF')
        self._nom_dest = ultimo('Destino', 'Nombre Destino')
        self._nom_cen = ultimo('Centro', 'Nombre Cedis')
        self._desc_mat = ultimo('Material', 'Desc. Material')

    def cond_exp(self, sf, centro, destino, material):
        r = self._l1.get(f'{sf}|{centro}|{destino}|{material}')
        if r and r['Cond. Expedición']:
            return r['Cond. Expedición'], 'Base 2026 (ruta exacta)'
        c = self._cond3.get(f'{sf}|{centro}|{material}')
        if c:
            return c, 'Base 2026 (sin destino)'
        c = self._cond_csd.get(f'{centro}|{destino}|{material}')
        if c:
            return c, 'Base propia (Centro+Destino+Material)'
        return '', ''

    def pv(self, sf, centro, destino, material):
        r = self._l1.get(f'{sf}|{centro}|{destino}|{material}')
        if r and pd.notna(r['PV']):
            return float(r['PV']), 'Base 2026 (ruta exacta)'
        p = self._pv3.get(f'{sf}|{centro}|{material}')
        if p is not None and pd.notna(p):
            return float(p), 'Base 2026 (sin destino)'
        p = self._pv_csd.get(f'{centro}|{destino}|{material}')
        if p is not None and pd.notna(p):
            return float(p), 'Base propia (Centro+Destino+Material)'
        p = self._pv_cm.get(f'{centro}|{material}')
        if p is not None and pd.notna(p):
            return float(p), 'Base propia (Centro+Material)'
        return None, ''

    def nombre_sf(self, sf):
        return self._nom_sf.get(sf, '')

    def nombre_destino(self, destino):
        return self._nom_dest.get(destino, '')

    def nombre_centro(self, centro):
        return self._nom_cen.get(centro, '')

    def desc_material(self, material):
        return self._desc_mat.get(material, '')


# ---------------------------------------------------------------------------------------------
# Calidad y Excel de la base
# ---------------------------------------------------------------------------------------------
def _conflictos(bd):
    v = bd[~bd['Es baja']]
    filas = []
    for campo in ('Cond. Expedición', 'PV', 'Nombre Destino'):
        w = v[v[campo].notna() & (v[campo] != '')]
        g = w.groupby('Llave')[campo].agg(lambda s: sorted({str(x) for x in s}))
        for llave, valores in g.items():
            if len(valores) > 1:
                usado = bd[(bd['Llave'] == llave) & bd['Registro usado en búsqueda']][campo]
                filas.append((llave, campo, ' | '.join(valores), str(usado.iloc[0]) if len(usado) else '',
                              int((w['Llave'] == llave).sum())))
    return filas


def resumen_calidad(bd, conf=None):
    v = bd[~bd['Es baja']]
    solo_baja = bd.groupby('Llave')['Es baja'].all()
    if conf is None:
        conf = _conflictos(bd)
    n_sf = bd[bd['Nombre SF'] != ''].groupby('Ship From')['Nombre SF'].nunique()
    n_dest = bd[bd['Nombre Destino'] != ''].groupby('Destino')['Nombre Destino'].nunique()
    fechas = bd['Fecha Procesamiento'].dropna()
    return [
        ('Registros en la base propia (todas las fuentes)', len(bd)),
    ] + [(f'   registros de: {f}', int(n)) for f, n in bd['Fuente'].value_counts().items()] + [
        ('Rutas distintas (SF+Centro+Destino+Material)', bd['Llave'].nunique()),
        ('Registros que son BAJA / CANCELADA (no se usan como fuente)', int(bd['Es baja'].sum())),
        ('Rutas cuyo único movimiento es BAJA / CANCELADA', int(solo_baja.sum())),
        ('Rutas disponibles para búsqueda', v['Llave'].nunique()),
        ('Registros sin modalidad', int((bd['Cond. Expedición'] == '').sum())),
        ('Registros con modalidad distinta de 1 / 2 / 4',
         int((~bd['Cond. Expedición'].isin(['', '1', '2', '4'])).sum())),
        ('Registros sin PV', int(bd['PV'].isna().sum())),
        ('Registros sin Fecha de Procesamiento', int(bd['Fecha Procesamiento'].isna().sum())),
        ('Rutas con modalidad distinta entre registros', len({c[0] for c in conf if c[1] == 'Cond. Expedición'})),
        ('Rutas con PV distinto entre registros', len({c[0] for c in conf if c[1] == 'PV'})),
        ('Rutas con Nombre Destino distinto entre registros', len({c[0] for c in conf if c[1] == 'Nombre Destino'})),
        ('Códigos de SF con más de un nombre', int((n_sf > 1).sum())),
        ('Códigos de destino con más de un nombre', int((n_dest > 1).sum())),
        ('Fecha de Procesamiento más antigua', fechas.min().strftime('%d/%m/%Y') if len(fechas) else ''),
        ('Fecha de Procesamiento más reciente', fechas.max().strftime('%d/%m/%Y') if len(fechas) else ''),
        ('Base generada el', datetime.now().strftime('%d/%m/%Y %H:%M')),
    ]


_NAVY = PatternFill('solid', fgColor='FFB24A24')      # terracota
_ZEBRA = PatternFill('solid', fgColor='FFF3F6FB')
_BORDE = Border(bottom=Side(style='thin', color='FFD9E1EC'))


def _escribir_hoja(wb, titulo, encabezados, filas, anchos=None, primera=False, fechas=(), congelar=True):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo
    ws.append(list(encabezados))
    for c in ws[1]:
        c.fill = _NAVY
        c.font = Font(name='Segoe UI', bold=True, color='FFFFFFFF', size=10)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for fila in filas:
        ws.append([None if (isinstance(x, float) and pd.isna(x)) or x is pd.NaT else x for x in fila])
    # Estilo celda por celda es MUY lento en hojas grandes (17 s con 5 mil filas): solo se aplica a las
    # hojas pequeñas; en las grandes queda el estilo por defecto y solo se da formato a las fechas.
    idx_fechas = [i + 1 for i, h in enumerate(encabezados) if h in fechas]
    pequena = ws.max_row <= 600
    if pequena or idx_fechas:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if pequena:
                    c.font = Font(name='Segoe UI', size=9)
                    c.border = _BORDE
                if c.column in idx_fechas and c.value is not None:
                    c.number_format = 'dd/mm/yyyy'
    for i, h in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(h, max(12, min(40, len(str(h)) + 4)))
    if congelar:
        ws.freeze_panes = 'A2'
    if len(encabezados) and ws.max_row > 1:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(encabezados))}{ws.max_row}'
    return ws


def guardar_bd_excel(bd, ruta):
    """Escribe un Excel con SOLO la hoja BD_Completa (una fila por registro, todas las fuentes)."""
    anchos = {'Fuente': 34, 'Llave': 34, 'Llave sin destino': 28, 'Llave sin SF': 26, 'Llave centro-material': 22,
              'Nombre Cedis': 24, 'Nombre SF': 30, 'Nombre Destino': 34, 'Desc. Material': 34, 'Observación': 36,
              'Ruta': 30}
    wb = Workbook()
    _escribir_hoja(wb, 'BD_Completa', COLUMNAS_BD, bd[COLUMNAS_BD].itertuples(index=False, name=None),
                   anchos=anchos, primera=True, fechas=('Fecha Procesamiento', 'Fecha Solicitud'))
    wb.save(ruta)
    return ruta


def resumen_texto(bd):
    """Una línea con lo más importante de la calidad de la base (para el log)."""
    conf = _conflictos(bd)
    return (f"{len(bd)} registros de {bd['Fuente'].nunique()} fuentes | rutas distintas: {bd['Llave'].nunique()} | "
            f"bajas: {int(bd['Es baja'].sum())} | rutas con datos contradictorios: {len({c[0] for c in conf})}")
